import docx
import matplotlib.pyplot as plt
import numpy as np
import pytest
from docx.oxml.ns import qn

from scripts import generate_report_docx


@pytest.fixture
def report_paths(monkeypatch, tmp_path):
    output = tmp_path / "report.docx"
    figures = tmp_path / "figures"
    figures.mkdir()
    monkeypatch.setattr(generate_report_docx, "DOC_OUT", output)
    monkeypatch.setattr(generate_report_docx, "EXP_DIR", figures)
    return output, figures


def test_build_report_without_optional_figures_contains_sections_and_tables(report_paths):
    output, _ = report_paths

    generate_report_docx.build_report()

    report = docx.Document(output)
    headings = [p.text for p in report.paragraphs if p.style.name == "Heading 1"]
    assert any(text.startswith("1. MỤC TIÊU") for text in headings)
    assert any(text.startswith("5. KẾT QUẢ") for text in headings)
    assert any(text.startswith("7. KẾT LUẬN") for text in headings)
    assert any(text.startswith("PHỤ LỤC") for text in headings)
    assert len(report.inline_shapes) == 0

    benchmark_table = next(
        table for table in report.tables
        if table.cell(0, 0).text == "Hàm" and table.cell(0, 1).text == "Dim"
    )
    assert len(benchmark_table.rows) == 19
    assert len(benchmark_table.columns) == 9
    assert {benchmark_table.cell(row, 0).text for row in range(1, 19)} == {
        "sphere", "schwefel_2_22", "rosenbrock", "rastrigin", "ackley", "griewank"
    }
    assert all(
        benchmark_table.cell(row, 8).text in {"DBO", "IDBO", "Hòa"}
        for row in range(1, 19)
    )


@pytest.mark.parametrize("figure_name,caption", [
    ("fig_convergence_dim10.png", "Hình F1: So sánh tốc độ hội tụ"),
    ("fig_boxplot_dim10.png", "Hình F2: Phân bố độ ổn định"),
    ("fig_diversity_dim10.png", "Hình F3: Sự biến thiên"),
])
def test_build_report_embeds_each_available_figure_independently(report_paths, figure_name, caption):
    output, figures = report_paths
    plt.imsave(figures / figure_name, np.zeros((2, 2, 3)))

    generate_report_docx.build_report()

    report = docx.Document(output)
    assert len(report.inline_shapes) == 1
    assert any(paragraph.text.startswith(caption) for paragraph in report.paragraphs)


def test_table_style_formats_header_and_alternating_rows():
    table = docx.Document().add_table(rows=3, cols=1)
    for row, text in enumerate(("Header", "First", "Second")):
        table.cell(row, 0).text = text

    generate_report_docx.style_table(table)

    assert table.cell(0, 0).paragraphs[0].runs[0].font.bold
    fills = [
        table.cell(row, 0)._tc.get_or_add_tcPr().find(qn("w:shd")).get(qn("w:fill"))
        for row in range(3)
    ]
    assert fills == [generate_report_docx.PRIMARY_HEX, generate_report_docx.BG_LIGHT_HEX, "FFFFFF"]
