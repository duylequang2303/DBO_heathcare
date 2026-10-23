"""Dựng slide báo cáo Tuần 11 (Chuẩn hóa thực nghiệm DBO/IDBO, Budget-Matched & Tích hợp Web UI).

Đề tài: CNTT-KLCN142 — Trường Đại học Công Thương TP.HCM (HUIT)
Thực hiện: Lê Quang Duy (Nhóm trưởng), Hồ Trung Cương, Đặng Nguyễn Minh Đăng.

Cấu trúc trình chiếu 10 slide chuẩn học thuật (16:9 Widescreen):
- Slide 1: Bìa báo cáo tổng quan Tuần 11
- Slide 2: Mục tiêu trọng tâm & Ma trận phân công Tuần 11
- Slide 3: Tích hợp Giao diện Web UI & API (Nghiệm thu Tuần 10)
- Slide 4: Chuẩn hóa thực nghiệm DBO vs IDBO (Ngưỡng 0.28 trên Tập độc lập)
- Slide 5: Phân tích Thống kê Paired & Kiểm định Wilcoxon Signed-Rank
- Slide 6: Phân tích Telemetry — Đóng góp của Tầng Ngoài vs Tầng Trong
- Slide 7: Kiểm soát Ngân sách Tính toán (Budget-Matched Benchmark)
- Slide 8: Kiến trúc Ngưỡng Kép (Dual Diversity Threshold)
- Slide 9: Đảm bảo Chất lượng & Kiểm thử Hồi quy (253/253 tests PASS)
- Slide 10: Tổng kết Tuần 11 & Kế hoạch Tuần 12
"""

from __future__ import annotations

import sys
from pathlib import Path
import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "docs" / "slides"
OUT_DIR.mkdir(parents=True, exist_ok=True)
OUT_PPTX = OUT_DIR / "Slide_Bao_Cao_Tuan_11_DBO_IDBO.pptx"

# --------------------------------------------------------------------------- #
# Bảng màu chuẩn học thuật (Academic Dark Navy / Teal / Coral)
# --------------------------------------------------------------------------- #
INDIGO = RGBColor(0x1B, 0x2A, 0x4A)       # Xanh Navy đậm chủ đạo
INDIGO2 = RGBColor(0x2D, 0x3E, 0x63)      # Xanh Navy phụ
AMBER = RGBColor(0xD9, 0x82, 0x2B)        # Cam Vàng điểm nhấn
TEAL = RGBColor(0x1D, 0x6F, 0x6F)         # Xanh Teal
GREEN = RGBColor(0x21, 0x7A, 0x4B)        # Xanh lá tích cực
RED = RGBColor(0xB8, 0x32, 0x28)          # Đỏ cảnh báo / điểm nghẽn
INK = RGBColor(0x1E, 0x24, 0x30)          # Đen chữ chính
MUTED = RGBColor(0x5A, 0x64, 0x78)        # Xám ghi chữ phụ
LIGHT = RGBColor(0xF5, 0xF7, 0xFB)        # Nền thẻ sáng
LINE = RGBColor(0xD3, 0xDB, 0xEA)         # Đường viền ngăn cách
WHITE = RGBColor(0xFF, 0xFF, 0xFF)        # Trắng tinh
CREAM = RGBColor(0xFD, 0xF6, 0xEB)        # Nền kem ghi chú
CARD_BORDER = RGBColor(0xE2, 0xE8, 0xF0)  # Viền thẻ trang nhã

FONT = "Calibri"
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def _add_run(paragraph, text=""):
    r = paragraph.add_run()
    if text:
        r.text = text
    return r


def _style_run(run, size, bold, color, italic=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color


def _add_header(slide, title_text, category_text):
    """Header chuẩn học thuật: Thẻ phân loại + Tiêu đề slide."""
    cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.35))
    tf_cat = cat_box.text_frame
    tf_cat.word_wrap = True
    tf_cat.margin_left = tf_cat.margin_top = tf_cat.margin_right = tf_cat.margin_bottom = 0
    p_cat = tf_cat.paragraphs[0]
    r_cat = _add_run(p_cat, category_text.upper())
    _style_run(r_cat, 11, True, AMBER)

    t_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.7), Inches(0.65))
    tf_t = t_box.text_frame
    tf_t.word_wrap = True
    tf_t.margin_left = tf_t.margin_top = tf_t.margin_right = tf_t.margin_bottom = 0
    p_t = tf_t.paragraphs[0]
    r_t = _add_run(p_t, title_text)
    _style_run(r_t, 22, True, INDIGO)

    # Đường kẻ phân cách
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.45), Inches(11.733), Inches(0.02))
    line.fill.solid()
    line.fill.fore_color.rgb = LINE
    line.line.color.rgb = LINE


def _add_card(slide, left, top, width, height, bg_color=LIGHT, border_color=CARD_BORDER):
    """Tạo khung chứa nội dung (Card Container)."""
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    if border_color:
        card.line.color.rgb = border_color
        card.line.width = Pt(1)
    else:
        card.line.fill.background()
    return card


# =========================================================================== #
# SLIDE 1: BÌA BÁO CÁO
# =========================================================================== #
def build_slide_1(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Nền Navy sang trọng
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = INDIGO
    bg.line.fill.background()

    # Dải màu cam trang trí
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.2), Inches(0.15), Inches(4.8))
    bar.fill.solid()
    bar.fill.fore_color.rgb = AMBER
    bar.line.fill.background()

    # Text Frame
    tb = slide.shapes.add_textbox(Inches(1.2), Inches(1.2), Inches(11.0), Inches(5.0))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

    p0 = tf.paragraphs[0]
    r0 = _add_run(p0, "TRƯỜNG ĐẠI HỌC CÔNG THƯƠNG TP. HỒ CHÍ MINH — KHOA CÔNG NGHỆ THÔNG TIN\nKHÓA LUẬN TỐT NGHIỆP | MÃ ĐỀ TÀI: CNTT-KLCN142\n")
    _style_run(r0, 13, True, AMBER)

    p1 = tf.add_paragraph()
    p1.space_before = Pt(14)
    r1 = _add_run(p1, "Tối Ưu Hóa Thực Đơn Dinh Dưỡng Cá Nhân Hóa\nBằng Thuật Toán DBO và IDBO")
    _style_run(r1, 30, True, WHITE)

    p2 = tf.add_paragraph()
    p2.space_before = Pt(14)
    r2 = _add_run(p2, "BÁO CÁO TIẾN ĐỘ TUẦN 11 — CHUẨN HÓA THỰC NGHIỆM ĐỘC LẬP,\nĐÁNH GIÁ BUDGET-MATCHED & TÍCH HỢP HỆ THỐNG WEB API")
    _style_run(r2, 16, True, RGBColor(0x6E, 0xC5, 0xC5))

    p3 = tf.add_paragraph()
    p3.space_before = Pt(28)
    r3 = _add_run(p3, "Sinh viên thực hiện:\n"
                      "• Lê Quang Duy (2001210356) — Nhóm trưởng (Thống kê, Báo cáo & Phân tích)\n"
                      "• Hồ Trung Cương (2001215682) — Phụ trách Budget-Matched & Dual Threshold\n"
                      "• Đặng Nguyễn Minh Đăng (2001230175) — Phụ trách Web/API & Trải nghiệm Người dùng\n\n"
                      "Thời gian: Tháng 10/2026 | Bộ môn: Kỹ thuật Dữ liệu & Trí tuệ Nhân tạo")
    _style_run(r3, 13, False, RGBColor(0xE2, 0xE8, 0xF0))


# =========================================================================== #
# SLIDE 2: MỤC TIÊU & MA TRẬN PHÂN CÔNG TUẦN 11
# =========================================================================== #
def build_slide_2(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_header(slide, "Mục Tiêu Trọng Tâm & Ma Trận Phân Công Tuần 11", "Kế Hoạch & Điều Phối")

    # Card 1: 4 Trọng tâm kỹ thuật
    _add_card(slide, Inches(0.8), Inches(1.65), Inches(5.6), Inches(5.3))
    tb1 = slide.shapes.add_textbox(Inches(1.1), Inches(1.85), Inches(5.0), Inches(4.9))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    p1 = tf1.paragraphs[0]
    r1 = _add_run(p1, "4 Trọng Tâm Nghiên Cứu Tuần 11")
    _style_run(r1, 16, True, INDIGO)

    items = [
        ("1. Nghiệm thu & Tích hợp Web UI/API:", "Hoàn thành giải quyết xung đột PR #28, đưa toàn bộ giao diện Dark Glassmorphism và API vào nhánh chính."),
        ("2. Chuẩn hóa thực nghiệm DBO vs IDBO:", "Tách bạch tập hiệu chỉnh (seed 8001-8002) và tập độc lập (seed 8003-8010). Giữ cố định ngưỡng 0.28."),
        ("3. Kiểm tra tính công bằng theo ngân sách:", "Thiết lập giới hạn số lần đánh giá hàm mục tiêu (max_evaluations) để trả lời: IDBO tốt hơn nhờ cơ chế hay nhờ thêm tài nguyên?"),
        ("4. Phân tách kiến trúc hai ngưỡng (Dual Threshold):", "Tách riêng outer_diversity_threshold (0.28) và inner_diversity_threshold để kiểm soát độc lập từng tầng.")
    ]
    for title, desc in items:
        p = tf1.add_paragraph()
        p.space_before = Pt(10)
        rt = _add_run(p, f"• {title}\n  ")
        _style_run(rt, 12.5, True, INDIGO2)
        rd = _add_run(p, desc)
        _style_run(rd, 11.5, False, INK)

    # Card 2: Phân công 3 thành viên
    _add_card(slide, Inches(6.8), Inches(1.65), Inches(5.7), Inches(5.3))
    tb2 = slide.shapes.add_textbox(Inches(7.1), Inches(1.85), Inches(5.1), Inches(4.9))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    p2 = tf2.paragraphs[0]
    r2 = _add_run(p2, "Ma Trận Trách Nhiệm Thành Viên")
    _style_run(r2, 16, True, INDIGO)

    roles = [
        ("Lê Quang Duy (Nhóm trưởng / Báo cáo & Thống kê)",
         "• Phân tích số liệu DBO vs IDBO tập độc lập (seed 8003-8010).\n"
         "• Thực hiện kiểm định Wilcoxon signed-rank, Sign test.\n"
         "• Giám sát tính trung thực, không phát ngôn vượt số liệu y tế.\n"
         "• Soạn thảo báo cáo Word và Slide thuyết trình Tuần 11."),
        ("Hồ Trung Cương (Budget-Matched & Dual Threshold)",
         "• Cài đặt cơ chế kiểm soát budget (max_evaluations) cho solver.\n"
         "• Tách 2 tham số outer/inner diversity threshold trong TwoTierConfig.\n"
         "• Bổ sung telemetry đo lường tài nguyên thực tế.\n"
         "• Xây dựng test suite kiểm tra budget dừng chính xác (PR #30)."),
        ("Đặng Nguyễn Minh Đăng (Tích hợp Web/API & Trải nghiệm)",
         "• Xây dựng giao diện Web responsive, theme glassmorphism.\n"
         "• Trực quan hóa dinh dưỡng (Radar Chart 6 trục, Progress bars dung sai).\n"
         "• Tích hợp 6 REST endpoints FastAPI, xử lý lỗi và schema validation (PR #28).")
    ]
    for name, task in roles:
        p = tf2.add_paragraph()
        p.space_before = Pt(8)
        rn = _add_run(p, f"👤 {name}\n")
        _style_run(rn, 12, True, TEAL)
        rk = _add_run(p, task)
        _style_run(rk, 11, False, INK)


# =========================================================================== #
# SLIDE 3: TÍCH HỢP GIAO DIỆN WEB UI & API (TUẦN 10)
# =========================================================================== #
def build_slide_3(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_header(slide, "Nghiệm Thu Tích Hợp Hệ Thống Web API & Giao Diện Người Dùng", "Hệ Thống Thực Nghiệm")

    # Card 1: Kiến trúc Web API
    _add_card(slide, Inches(0.8), Inches(1.65), Inches(5.6), Inches(5.3))
    tb1 = slide.shapes.add_textbox(Inches(1.1), Inches(1.85), Inches(5.0), Inches(4.9))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    p1 = tf1.paragraphs[0]
    r1 = _add_run(p1, "Kiến Trúc Kỹ Thuật (FastAPI Backend)")
    _style_run(r1, 16, True, INDIGO)

    features = [
        ("Thiết kế Single-Port (Một Cổng Phục Vụ):", "FastAPI phục vụ đồng thời cả REST API (/api/*) và tài nguyên tĩnh (/static) trên cổng 8000. Loại bỏ hoàn toàn sự cố CORS và proxy trung gian."),
        ("6 Endpoints Chuẩn Hóa:", "GET /health, GET /profiles (P1/P2/P3), POST /targets (tính BMR/TDEE tức thời), POST /menu/generate (chạy IDBO 2 tầng), GET /foods, POST /menu/swap (stub 501 minh bạch)."),
        ("Xử Lý Phi Nghẽn (Non-blocking):", "Sử dụng asyncio.to_thread để bao bọc quá trình tối ưu hóa 2 tầng, giữ cho Event Loop của FastAPI luôn phản hồi trơn tru."),
        ("Tối Ưu Hóa Bộ Nhớ (Singleton DB):", "Cơ sở dữ liệu 15.929 thực phẩm được nạp vào RAM một lần duy nhất tại sự kiện startup.")
    ]
    for t, d in features:
        p = tf1.add_paragraph()
        p.space_before = Pt(8)
        rt = _add_run(p, f"✓ {t}\n  ")
        _style_run(rt, 12, True, INDIGO2)
        rd = _add_run(p, d)
        _style_run(rd, 11, False, INK)

    # Card 2: Giao diện trực quan hóa
    _add_card(slide, Inches(6.8), Inches(1.65), Inches(5.7), Inches(5.3))
    tb2 = slide.shapes.add_textbox(Inches(7.1), Inches(1.85), Inches(5.1), Inches(4.9))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    p2 = tf2.paragraphs[0]
    r2 = _add_run(p2, "Trực Quan Hóa & Trải Nghiệm Người Dùng (UI)")
    _style_run(r2, 16, True, INDIGO)

    ui_points = [
        ("Thiết Kế Hiện Đại (Dark Glassmorphism):", "Giao diện tông tối sang trọng, không phụ thuộc framework CSS nặng, responsive linh hoạt trên máy tính và thiết bị di động."),
        ("Tuân Thủ Quy Chuẩn Khẩu Phần (AGENTS.md):", "Bảng thực đơn hiển thị minh bạch 2 cột: Gram liên tục gốc và Gram làm tròn bước 5g (portion_g_rounded). Dinh dưỡng được tính toán lại chính xác trên gram làm tròn."),
        ("Thanh Tiến Độ Dung Sai (Progress Bars):", "Chỉ báo 3 trạng thái màu sắc theo dải dung sai nghiêm ngặt: Xanh lục (đạt chuẩn ±10-20%), Vàng (cận biên), Đỏ (vượt ngưỡng hoặc thiếu hụt)."),
        ("Biểu Đồ Radar Đa Chiều (Chart.js):", "Trực quan hóa đồng thời 6 chỉ số: Năng lượng, Protein, Carbohydrate, Lipid, Chất xơ và Natri so với mục tiêu tham chiếu (100%).")
    ]
    for t, d in ui_points:
        p = tf2.add_paragraph()
        p.space_before = Pt(8)
        rt = _add_run(p, f"★ {t}\n  ")
        _style_run(rt, 12, True, TEAL)
        rd = _add_run(p, d)
        _style_run(rd, 11, False, INK)


# =========================================================================== #
# SLIDE 4: CHUẨN HÓA THỰC NGHIỆM DBO vs IDBO (TẬP ĐỘC LẬP)
# =========================================================================== #
def build_slide_4(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_header(slide, "Chuẩn Hóa Dữ Liệu Thực Nghiệm DBO vs IDBO (Ngưỡng 0.28)", "Đánh Giá Thực Nghiệm")

    # Card bảng số liệu
    _add_card(slide, Inches(0.8), Inches(1.65), Inches(11.733), Inches(5.3))

    tb = slide.shapes.add_textbox(Inches(1.1), Inches(1.8), Inches(11.1), Inches(0.5))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    r = _add_run(p, "Bảng Đối Chứng Chi Tiết Từng Hạt Giống (Tập Đánh Giá Độc Lập Seed 8003–8010)")
    _style_run(r, 15, True, INDIGO)

    # Thêm bảng
    rows, cols = 10, 8
    table_shape = slide.shapes.add_table(rows, cols, Inches(1.1), Inches(2.35), Inches(11.133), Inches(3.4))
    tbl = table_shape.table

    # Đặt độ rộng cột
    tbl.columns[0].width = Inches(1.1)   # Seed
    tbl.columns[1].width = Inches(1.5)   # DBO Fit
    tbl.columns[2].width = Inches(1.5)   # IDBO Fit
    tbl.columns[3].width = Inches(1.5)   # Diff
    tbl.columns[4].width = Inches(1.4)   # DBO Time
    tbl.columns[5].width = Inches(1.4)   # IDBO Time
    tbl.columns[6].width = Inches(1.4)   # Pert/Rest
    tbl.columns[7].width = Inches(1.333) # Outcome

    headers = ["Hạt giống", "DBO Fitness", "IDBO Fitness", "Chênh lệch", "Thời gian DBO", "Thời gian IDBO", "Đột biến/Restart", "Kết quả"]
    for j, h in enumerate(headers):
        cell = tbl.cell(0, j)
        cell.fill.solid()
        cell.fill.fore_color.rgb = INDIGO
        p = cell.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = _add_run(p, h)
        _style_run(run, 11, True, WHITE)

    data = [
        ("Seed 8003", "99.0691", "97.9416", "-1.1275", "52.4s", "57.3s", "28 / 3", "DBO Thắng", RED),
        ("Seed 8004", "97.2465", "99.1396", "+1.8931", "48.8s", "57.6s", "4 / 0", "IDBO Thắng", GREEN),
        ("Seed 8005", "99.1117", "99.7556", "+0.6439", "53.8s", "65.1s", "69 / 7", "IDBO Thắng", GREEN),
        ("Seed 8006", "99.3151", "98.9199", "-0.3952", "53.7s", "58.5s", "44 / 5", "DBO Thắng", RED),
        ("Seed 8007", "99.5862", "99.6609", "+0.0747", "55.8s", "64.7s", "0 / 0", "IDBO Thắng", GREEN),
        ("Seed 8008", "99.4844", "98.4347", "-1.0497", "54.5s", "57.0s", "0 / 0", "DBO Thắng", RED),
        ("Seed 8009", "98.5132", "99.6530", "+1.1398", "37.6s", "54.1s", "58 / 6", "IDBO Thắng", GREEN),
        ("Seed 8010", "99.6185", "99.7668", "+0.1483", "38.6s", "38.6s", "33 / 2", "IDBO Thắng", GREEN),
        ("Trung bình", "98.9931", "99.1590", "+0.1659", "49.39s", "56.61s", "29.5 / 2.9", "IDBO: 5 Thắng (62.5%)", INDIGO2)
    ]

    for i, row in enumerate(data, start=1):
        for j, val in enumerate(row[:8]):
            cell = tbl.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = LIGHT if i % 2 == 1 else WHITE
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            run = _add_run(p, val)
            color = row[8] if j in (3, 7) else INK
            bold = True if i == 9 or j in (0, 3, 7) else False
            _style_run(run, 10.5, bold, color)

    # Chú thích học thuật ở dưới
    tb_note = slide.shapes.add_textbox(Inches(1.1), Inches(5.9), Inches(11.1), Inches(0.8))
    tf_n = tb_note.text_frame
    tf_n.word_wrap = True
    pn = tf_n.paragraphs[0]
    rn = _add_run(pn, "(*) Nguyên tắc phân định dữ liệu: ")
    _style_run(rn, 11, True, AMBER)
    rn2 = _add_run(pn, "Tập hiệu chỉnh (Calibration seed 8001–8002: IDBO thắng 1, DBO thắng 1) chỉ dùng để khảo sát ngưỡng, tuyệt đối không gộp vào tập đánh giá độc lập để tránh thiên kiến chọn lọc (Selection Bias).")
    _style_run(rn2, 11, False, MUTED)


# =========================================================================== #
# SLIDE 5: PHÂN TÍCH THỐNG KÊ & KIỂM ĐỊNH WILCOXON
# =========================================================================== #
def build_slide_5(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_header(slide, "Phân Tích Thống Kê Đối Chứng & Kiểm Định Ý Nghĩa", "Phân Tích Khoa Học")

    # Card 1: Kết quả kiểm định thống kê
    _add_card(slide, Inches(0.8), Inches(1.65), Inches(5.6), Inches(5.3))
    tb1 = slide.shapes.add_textbox(Inches(1.1), Inches(1.85), Inches(5.0), Inches(4.9))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    p1 = tf1.paragraphs[0]
    r1 = _add_run(p1, "Kiểm Định Wilcoxon Signed-Rank Test")
    _style_run(r1, 16, True, INDIGO)

    stats_info = [
        ("Mẫu thử nghiệm cặp (Paired Samples):", "N = 8 hạt giống độc lập (seed 8003–8010)."),
        ("Tổng hạng dương (IDBO thắng):", "W+ = 22.0 (tổng hạng của 5 hạt giống có diff > 0)."),
        ("Tổng hạng âm (DBO thắng):", "W- = 14.0 (tổng hạng của 3 hạt giống có diff < 0)."),
        ("Trị thống kê kiểm định:", "W = min(W+, W-) = 14.0."),
        ("Giá trị tới hạn (Critical Value):", "Với n = 8, mức ý nghĩa alpha = 0.05 (2 phía), ngưỡng bác bỏ giả thuyết H0 là W <= 3."),
        ("Kết luận kiểm định thống kê:", "Vì W = 14.0 > 3 (p > 0.05), chưa đủ bằng chứng thống kê để khẳng định IDBO vượt trội hơn DBO một cách có ý nghĩa.")
    ]
    for t, d in stats_info:
        p = tf1.add_paragraph()
        p.space_before = Pt(6)
        rt = _add_run(p, f"• {t} ")
        _style_run(rt, 12, True, INDIGO2)
        rd = _add_run(p, d)
        _style_run(rd, 11.5, False, INK)

    # Card 2: Đánh đổi tài nguyên tính toán
    _add_card(slide, Inches(6.8), Inches(1.65), Inches(5.7), Inches(5.3))
    tb2 = slide.shapes.add_textbox(Inches(7.1), Inches(1.85), Inches(5.1), Inches(4.9))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    p2 = tf2.paragraphs[0]
    r2 = _add_run(p2, "Đánh Đổi Chi Phí Tính Toán (Trade-off)")
    _style_run(r2, 16, True, INDIGO)

    tradeoffs = [
        ("Chênh lệch điểm số trung bình:", "IDBO đạt 99.1590 so với 98.9931 của DBO (cải thiện nhẹ +0.1659 điểm, trung vị tăng +0.1115)."),
        ("Số lượt tính hàm mục tiêu (Evaluations):", "DBO cố định ở 263.870 lượt; IDBO tiêu tốn trung bình 313.005 lượt (tăng +18.6%, tối đa lên tới 337.808 lượt ở seed 8005)."),
        ("Thời gian thực thi (Runtime):", "DBO trung bình 49.39s; IDBO trung bình 56.61s (tăng +14.6% thời gian xử lý do chi phí tính toán đa dạng và restart)."),
        ("Tuyên bố khoa học trung thực:", "IDBO mang lại hành vi tìm kiếm đa dạng và có xu hướng cải thiện nhẹ, nhưng đi kèm chi phí tài nguyên cao hơn. Cần kiểm chứng trong điều kiện cố định ngân sách.")
    ]
    for t, d in tradeoffs:
        p = tf2.add_paragraph()
        p.space_before = Pt(6)
        rt = _add_run(p, f"▲ {t} ")
        _style_run(rt, 12, True, AMBER)
        rd = _add_run(p, d)
        _style_run(rd, 11.5, False, INK)


# =========================================================================== #
# SLIDE 6: PHÂN TÍCH TELEMETRY TẦNG NGOÀI VS TẦNG TRONG
# =========================================================================== #
def build_slide_6(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_header(slide, "Giải Mã Telemetry — Đóng Góp Của Tầng Ngoài và Tầng Trong", "Cơ Chế IDBO")

    _add_card(slide, Inches(0.8), Inches(1.65), Inches(5.6), Inches(5.3))
    tb1 = slide.shapes.add_textbox(Inches(1.1), Inches(1.85), Inches(5.0), Inches(4.9))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    p1 = tf1.paragraphs[0]
    r1 = _add_run(p1, "Hoạt Động Tầng Ngoài (Discrete Outer Solver)")
    _style_run(r1, 16, True, INDIGO)

    outer_pts = [
        ("Mục tiêu tầng ngoài:", "Tìm kiếm cấu trúc tổ hợp món ăn (food_ids) tối ưu cho các bữa ăn trong ngày."),
        ("Số lần đột biến (Perturbations):", "Trung bình 29.5 lần (dao động từ 0 đến 69 lần). Kích hoạt mạnh nhất tại seed 8005 (69 lần) và seed 8009 (58 lần)."),
        ("Số lần khởi động lại (Restarts):", "Trung bình 2.9 lần khi quần thể rơi vào tình trạng đình trệ nghiêm trọng."),
        ("Duy trì độ đa dạng quần thể:", "Độ đa dạng trung bình của IDBO tầng ngoài đạt 0.283, giúp tránh hội tụ sớm vào các tổ hợp món cục bộ.")
    ]
    for t, d in outer_pts:
        p = tf1.add_paragraph()
        p.space_before = Pt(8)
        rt = _add_run(p, f"• {t}\n  ")
        _style_run(rt, 12, True, INDIGO2)
        rd = _add_run(p, d)
        _style_run(rd, 11, False, INK)

    _add_card(slide, Inches(6.8), Inches(1.65), Inches(5.7), Inches(5.3))
    tb2 = slide.shapes.add_textbox(Inches(7.1), Inches(1.85), Inches(5.1), Inches(4.9))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    p2 = tf2.paragraphs[0]
    r2 = _add_run(p2, "Hoạt Động Tầng Trong (Continuous Inner Solver)")
    _style_run(r2, 16, True, INDIGO)

    inner_pts = [
        ("Mục tiêu tầng trong:", "Tối ưu hóa liên tục khối lượng khẩu phần gram cho từng món ăn nhằm bám sát chuẩn dinh dưỡng."),
        ("Cường độ kích hoạt cực lớn:", "Trung bình ~16.000 lần đột biến và 103 lần restart trên mỗi lần chạy hoàn chỉnh."),
        ("Hiện tượng Seed 8007 & 8008:", "Tại 2 seed này, tầng ngoài hoàn toàn không kích hoạt (0 đột biến, 0 restart) nhưng kết quả IDBO vẫn khác DBO (8007 thắng, 8008 thua) nhờ toàn bộ sự dịch chuyển diễn ra ở tầng trong."),
        ("Kết luận kiến trúc:", "Sự cải thiện của IDBO là nỗ lực kết hợp của cả hai tầng, trong đó tầng trong đóng vai trò tinh chỉnh vi mô quyết định.")
    ]
    for t, d in inner_pts:
        p = tf2.add_paragraph()
        p.space_before = Pt(8)
        rt = _add_run(p, f"★ {t}\n  ")
        _style_run(rt, 12, True, TEAL)
        rd = _add_run(p, d)
        _style_run(rd, 11, False, INK)


# =========================================================================== #
# SLIDE 7: KIỂM SOÁT NGÂN SÁCH (BUDGET-MATCHED BENCHMARK)
# =========================================================================== #
def build_slide_7(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_header(slide, "Kiểm Soát Tính Công Bằng Theo Ngân Sách Tính Toán (Budget-Matched)", "Thiết Kế Thử Nghiệm Mới")

    _add_card(slide, Inches(0.8), Inches(1.65), Inches(5.6), Inches(5.3))
    tb1 = slide.shapes.add_textbox(Inches(1.1), Inches(1.85), Inches(5.0), Inches(4.9))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    p1 = tf1.paragraphs[0]
    r1 = _add_run(p1, "Vấn Đề Khoa Học Đặt Ra")
    _style_run(r1, 16, True, INDIGO)

    q_pts = [
        ("Câu hỏi cốt lõi:", "IDBO có điểm số tốt hơn ở 5/8 seed là nhờ cơ chế tìm kiếm thông minh hơn hay chỉ đơn thuần vì được sử dụng nhiều lượt tính toán hơn (+18.6% evaluations)?"),
        ("Sự bất bình đẳng về tài nguyên:", "Khi cùng chạy 100 vòng lặp ngoài và 50 vòng lặp trong, DBO chỉ tốn 263k evaluations, trong khi IDBO tốn tới 313k evaluations do các bước nhảy đa dạng."),
        ("Giải pháp Budget-Matched:", "Đặt một giới hạn trần cố định max_evaluations cho cả DBO và IDBO. Thuật toán nào chạm trần trước sẽ phải dừng ngay lập tức."),
        ("Tiêu chí đánh giá mới:", "So sánh trực tiếp chất lượng thực đơn khi cả hai thuật toán tiêu thụ cùng một mức tài nguyên máy tính.")
    ]
    for t, d in q_pts:
        p = tf1.add_paragraph()
        p.space_before = Pt(8)
        rt = _add_run(p, f"? {t}\n  ")
        _style_run(rt, 12, True, RED)
        rd = _add_run(p, d)
        _style_run(rd, 11, False, INK)

    _add_card(slide, Inches(6.8), Inches(1.65), Inches(5.7), Inches(5.3))
    tb2 = slide.shapes.add_textbox(Inches(7.1), Inches(1.85), Inches(5.1), Inches(4.9))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    p2 = tf2.paragraphs[0]
    r2 = _add_run(p2, "Triển Khai & Kết Quả Thử Nghiệm Kỹ Thuật (PR #30)")
    _style_run(r2, 16, True, INDIGO)

    impl_pts = [
        ("Cơ chế Budget Enforcement Nghiêm Ngặt:", "Cài đặt theo dõi ngân sách còn lại (_budget_remaining) xuyên suốt từ tầng ngoài, các đợt đánh giá tầng trong đến đợt tối ưu cuối cùng (Final Inner)."),
        ("Khóa Cấu Hình & Commit Provenance:", "Script experiment_week11_budget.py tự động khóa mã băm cấu hình (config_hash) và Git commit hash vào config_locked.json."),
        ("Lưu trữ gia tăng (Incremental Flush):", "Lưu ngay kết quả từng seed (runs/8001_dbo.json) để đảm bảo không mất dữ liệu nếu bị ngắt."),
        ("Kết quả Smoke Test (Ngân sách 150 evaluations):", "Cả DBO và IDBO trên seed 8001 và 8002 đều dừng chuẩn xác ở 150/150 evaluations, không vượt ngân sách (budget_exceeded = False).")
    ]
    for t, d in impl_pts:
        p = tf2.add_paragraph()
        p.space_before = Pt(8)
        rt = _add_run(p, f"✓ {t}\n  ")
        _style_run(rt, 12, True, GREEN)
        rd = _add_run(p, d)
        _style_run(rd, 11, False, INK)


# =========================================================================== #
# SLIDE 8: KIẾN TRÚC NGƯỠNG KÉP (DUAL DIVERSITY THRESHOLD)
# =========================================================================== #
def build_slide_8(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_header(slide, "Kiến Trúc Phân Tách Ngưỡng Đa Dạng Kép (Dual Threshold)", "Hoàn Thiện Kiến Trúc")

    _add_card(slide, Inches(0.8), Inches(1.65), Inches(5.6), Inches(5.3))
    tb1 = slide.shapes.add_textbox(Inches(1.1), Inches(1.85), Inches(5.0), Inches(4.9))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    p1 = tf1.paragraphs[0]
    r1 = _add_run(p1, "Hạn Chế Của Kiến Trúc Cũ (Một Ngưỡng Dùng Chung)")
    _style_run(r1, 16, True, INDIGO)

    old_pts = [
        ("Bản chất không gian khác biệt:", "Tầng ngoài hoạt động trên không gian nhị phân/rời rạc (chọn món), trong khi tầng trong hoạt động trên không gian liên tục (chọn số gram khẩu phần)."),
        ("Độ đo đa dạng khác nhau:", "Tầng ngoài đo khoảng cách Hamming/Jaccard giữa các tổ hợp món; tầng trong đo độ lệch chuẩn vị trí của các cá thể liên tục."),
        ("Bất cập khi dùng chung 0.28:", "Ép cùng một ngưỡng 0.28 cho cả hai tầng là giải pháp tạm thời ở Tuần 10, chưa phản ánh đúng đặc tính phân bố nghiệm ở tầng trong."),
        ("Yêu cầu cải tiến Tuần 11:", "Tách độc lập hai tham số điều khiển ngưỡng để tối ưu hóa chuyên biệt.")
    ]
    for t, d in old_pts:
        p = tf1.add_paragraph()
        p.space_before = Pt(8)
        rt = _add_run(p, f"• {t}\n  ")
        _style_run(rt, 12, True, RED)
        rd = _add_run(p, d)
        _style_run(rd, 11, False, INK)

    _add_card(slide, Inches(6.8), Inches(1.65), Inches(5.7), Inches(5.3))
    tb2 = slide.shapes.add_textbox(Inches(7.1), Inches(1.85), Inches(5.1), Inches(4.9))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    p2 = tf2.paragraphs[0]
    r2 = _add_run(p2, "Cấu Trúc TwoTierConfig Tách Biệt Mới")
    _style_run(r2, 16, True, INDIGO)

    new_pts = [
        ("outer_diversity_threshold = 0.28:", "Được cố định cho tầng ngoài sau khi đã hiệu chỉnh và kiểm nghiệm độc lập."),
        ("inner_diversity_threshold = None (Mặc định):", "Theo nguyên tắc AGENTS.md, nhóm chưa tự ý gán một giá trị cảm tính khi chưa đo lường thực tế."),
        ("Bổ sung trường đo lường thực nghiệm:", "Thu thập hai trường telemetry mới: inner_diversity_min_observed và inner_diversity_mean_observed trong từng lần chạy."),
        ("Độ phủ kiểm thử:", "Đã bổ sung bộ unit tests kiểm tra việc truyền độc lập hai ngưỡng vào từng solver con mà không gây xung đột.")
    ]
    for t, d in new_pts:
        p = tf2.add_paragraph()
        p.space_before = Pt(8)
        rt = _add_run(p, f"★ {t}\n  ")
        _style_run(rt, 12, True, TEAL)
        rd = _add_run(p, d)
        _style_run(rd, 11, False, INK)


# =========================================================================== #
# SLIDE 9: ĐẢM BẢO CHẤT LƯỢNG MÃ NGUỒN & KIỂM THỬ HỒI QUY
# =========================================================================== #
def build_slide_9(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_header(slide, "Đảm Bảo Chất Lượng Phần Mềm & Kiểm Thử Hồi Quy", "Kỹ Thuật Phần Mềm")

    # Card 1: Số lượng test
    _add_card(slide, Inches(0.8), Inches(1.65), Inches(3.6), Inches(5.3))
    tb1 = slide.shapes.add_textbox(Inches(1.0), Inches(1.85), Inches(3.2), Inches(4.9))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    p1 = tf1.paragraphs[0]
    r1 = _add_run(p1, "100% Tests Pass")
    _style_run(r1, 16, True, INDIGO)

    p_stat = tf1.add_paragraph()
    p_stat.space_before = Pt(14)
    r_stat = _add_run(p_stat, "253 / 253\nPASSED")
    _style_run(r_stat, 28, True, GREEN)

    p_time = tf1.add_paragraph()
    p_time.space_before = Pt(6)
    r_time = _add_run(p_time, "Thời gian thực thi: 22.59s\npytest v9.1.1 | Python 3.12")
    _style_run(r_time, 12, False, MUTED)

    p_cov = tf1.add_paragraph()
    p_cov.space_before = Pt(14)
    r_cov = _add_run(p_cov, "Phân bổ kiểm thử:\n"
                            "• Core Solver & Behaviors: 82\n"
                            "• Constraints & Menus: 45\n"
                            "• Culinary Realism & Roles: 36\n"
                            "• Data Loader & Schemas: 42\n"
                            "• Budget Enforcement (W11): 27\n"
                            "• FastAPI Routes & Profile: 21")
    _style_run(r_cov, 11.5, False, INK)

    # Card 2: Tuân thủ bất biến nghiệp vụ
    _add_card(slide, Inches(4.7), Inches(1.65), Inches(7.8), Inches(5.3))
    tb2 = slide.shapes.add_textbox(Inches(5.0), Inches(1.85), Inches(7.2), Inches(4.9))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    p2 = tf2.paragraphs[0]
    r2 = _add_run(p2, "Bảo Vệ Các Bất Biến Nghiệp Vụ (Theo AGENTS.md)")
    _style_run(r2, 16, True, INDIGO)

    invariants = [
        ("Deterministic Decoding (Tái Lập Tuyệt Đối):", "Mọi hàm ngẫu nhiên đều nhận hạt giống (seed) tường minh và dùng RNG cục bộ (np.random.default_rng(seed)). Cùng seed luôn tạo ra cùng một thực đơn chính xác."),
        ("Quy Trình Làm Tròn Khẩu Phần 5g:", "Nghiệm liên tục -> Làm tròn khẩu phần 5g -> Giải mã thực đơn -> Tính toán lại dinh dưỡng -> Thẩm định tính khả thi (validate_menu). Không báo cáo fitness của nghiệm liên tục nhưng hiển thị thực đơn đã làm tròn khác."),
        ("Minh Bạch Dữ Liệu Đường & Dinh Dưỡng:", "Không tuyên bố dữ liệu đường đo thực tế nếu dùng công thức ước tính (sugar_g_estimated). Không đưa ra tuyên bố y tế lâm sàng hoặc cam kết điều trị bệnh."),
        ("Loại Bỏ Thực Phẩm Thô Khỏi Output:", "Bộ lọc Role-based bounds loại trừ hoàn toàn các loại bột sống, shortening, phụ gia khỏi candidate pool trước khi đưa vào giao diện đề xuất cho người dùng.")
    ]
    for t, d in invariants:
        p = tf2.add_paragraph()
        p.space_before = Pt(8)
        rt = _add_run(p, f"🛡 {t}\n  ")
        _style_run(rt, 12, True, INDIGO2)
        rd = _add_run(p, d)
        _style_run(rd, 11, False, INK)


# =========================================================================== #
# SLIDE 10: TỔNG KẾT TUẦN 11 & KẾ HOẠCH TUẦN 12
# =========================================================================== #
def build_slide_10(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_header(slide, "Tổng Kết Tuần 11 & Kế Hoạch Nghiên Cứu Tuần 12", "Kết Luận & Định Hướng")

    _add_card(slide, Inches(0.8), Inches(1.65), Inches(5.6), Inches(5.3))
    tb1 = slide.shapes.add_textbox(Inches(1.1), Inches(1.85), Inches(5.0), Inches(4.9))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    p1 = tf1.paragraphs[0]
    r1 = _add_run(p1, "Thành Quả Đạt Được Trong Tuần 11")
    _style_run(r1, 16, True, INDIGO)

    achievements = [
        ("Tích hợp hoàn tất hệ thống Web:", "Giải quyết conflict PR #28, đưa FastAPI backend và giao diện Dark Glassmorphism vào main, 253 tests pass."),
        ("Chuẩn hóa thực nghiệm DBO/IDBO:", "Phân tách rõ ràng tập calibration vs independent eval. Xác nhận IDBO thắng 5/8 seed nhưng tốn thêm 18.6% evaluations và kiểm định Wilcoxon chưa cho thấy sự vượt trội có ý nghĩa thống kê."),
        ("Hoàn thiện cơ chế Budget-Matched (PR #30):", "Cài đặt thành công cơ chế giới hạn evaluations chung, khóa cấu hình và commit Git tự động, smoke test 150 evaluations thành công."),
        ("Kiến trúc Ngưỡng kép (Dual Threshold):", "Tách riêng biệt outer và inner threshold, bổ sung telemetry đo độ đa dạng tầng trong.")
    ]
    for t, d in achievements:
        p = tf1.add_paragraph()
        p.space_before = Pt(8)
        rt = _add_run(p, f"✓ {t}\n  ")
        _style_run(rt, 12, True, GREEN)
        rd = _add_run(p, d)
        _style_run(rd, 11, False, INK)

    _add_card(slide, Inches(6.8), Inches(1.65), Inches(5.7), Inches(5.3))
    tb2 = slide.shapes.add_textbox(Inches(7.1), Inches(1.85), Inches(5.1), Inches(4.9))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    p2 = tf2.paragraphs[0]
    r2 = _add_run(p2, "Kế Hoạch Trọng Tâm Tuần 12")
    _style_run(r2, 16, True, INDIGO)

    plans = [
        ("1. Chạy Benchmark Budget-Matched Đầy Đủ:", "Chạy benchmark toàn diện trên tập seed độc lập với ngân sách chung (ví dụ 260.000 evaluations) để trả lời dứt điểm câu hỏi về tính công bằng tài nguyên."),
        ("2. Hiệu Chỉnh Ngưỡng Đa Dạng Tầng Trong:", "Sử dụng dữ liệu quan sát inner_diversity_mean để thiết lập và thử nghiệm ngưỡng inner_diversity_threshold phù hợp."),
        ("3. Kiểm Thử Đầu Cuối (E2E Web Test):", "Thực hiện tự động hóa kiểm thử luồng Người dùng -> UI Web -> API -> Two-Tier Solver -> Hiển thị thực đơn."),
        ("4. Hoàn Thiện Bản Thảo Khóa Luận Tốt Nghiệp:", "Tổng hợp toàn bộ các kết quả thực nghiệm từ Tuần 1 đến Tuần 12 vào báo cáo Word chính thức của đề tài.")
    ]
    for t, d in plans:
        p = tf2.add_paragraph()
        p.space_before = Pt(8)
        rt = _add_run(p, f"★ {t}\n  ")
        _style_run(rt, 12, True, AMBER)
        rd = _add_run(p, d)
        _style_run(rd, 11, False, INK)


def main():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    print("Building Week 11 Presentation Slides...")
    build_slide_1(prs)
    build_slide_2(prs)
    build_slide_3(prs)
    build_slide_4(prs)
    build_slide_5(prs)
    build_slide_6(prs)
    build_slide_7(prs)
    build_slide_8(prs)
    build_slide_9(prs)
    build_slide_10(prs)

    prs.save(str(OUT_PPTX))
    print(f"[SUCCESS] Slide presentation saved to: {OUT_PPTX}")
    print(f"[METRICS] Total slides: {len(prs.slides)}")


if __name__ == "__main__":
    main()
