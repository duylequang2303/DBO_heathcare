"""Dựng slide báo cáo tổng hợp Tuần 9 – 10 – 11 (DBO / IDBO Healthcare Menu Optimization).

Đề tài: CNTT-KLCN142 — Trường Đại học Công Thương TP.HCM (HUIT)
Thực hiện: Lê Quang Duy (Nhóm trưởng), Hồ Trung Cương, Đặng Nguyễn Minh Đăng.

Chuỗi tiến độ gộp 3 tuần trọng tâm (10 slide chuẩn học thuật 16:9, đầy đủ hình ảnh dẫn chứng thực tế):
- Slide 1: Bìa báo cáo tổng quan Tuần 9 – 10 – 11
- Slide 2: [Tuần 9] Lập thực đơn chu kỳ 7 ngày & Cơ chế Diversity Penalty (Kèm hình Heatmap W9-F1)
- Slide 3: [Tuần 9] Đối chuẩn trên 3 hồ sơ bệnh lý (P1, P2, P3) & Trải nghiệm UX Option A/B/C (Kèm hình W9-F2 & W9-F3)
- Slide 4: [Tuần 10] Nghiệm thu Hệ thống Web API & Giao diện Người dùng Web UI (Kèm ảnh chụp màn hình thực tế)
- Slide 5: [Tuần 11] Thực nghiệm đối đầu DBO vs IDBO trên Tập độc lập Seed 8003–8010 (Kèm hình W11-F1)
- Slide 6: [Tuần 11] Đánh đổi chi phí tính toán: Lượt Evaluations & Thời gian chạy (Kèm hình W11-F2)
- Slide 7: [Tuần 11] Giải mã Telemetry: Hoạt động Tầng ngoài vs Tầng trong (Kèm hình W11-F3)
- Slide 8: [Tuần 11] Kiểm soát ngân sách tính toán chung (Budget-Matched Benchmark, Kèm hình W11-F4)
- Slide 9: [Tuần 11] Kiến trúc Ngưỡng kép (Dual Threshold) & Nghiệm thu chất lượng 253/253 tests PASS
- Slide 10: Tổng kết giai đoạn Tuần 9–10–11 & Định hướng nghiên cứu Tuần 12
"""

from __future__ import annotations

import sys
from pathlib import Path
from typing import Tuple
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "docs" / "slides"
OUT_DIR.mkdir(parents=True, exist_ok=True)
OUT_PPTX = OUT_DIR / "Slide_Bao_Cao_Tuan_9_10_11_DBO_IDBO.pptx"
OUT_PPTX_W11 = OUT_DIR / "Slide_Bao_Cao_Tuan_11_DBO_IDBO.pptx"

W9_DIR = ROOT / "experiments" / "week9"
W11_DIR = ROOT / "experiments" / "ablation_dbo_idbo"

# --------------------------------------------------------------------------- #
# Bảng màu chuẩn học thuật (Academic Dark Navy / Teal / Coral)
# --------------------------------------------------------------------------- #
INDIGO = RGBColor(0x1B, 0x2A, 0x4A)       # Xanh Navy đậm chủ đạo
INDIGO2 = RGBColor(0x2D, 0x3E, 0x63)      # Xanh Navy phụ
AMBER = RGBColor(0xD9, 0x82, 0x2B)        # Cam Vàng điểm nhấn
TEAL = RGBColor(0x1D, 0x6F, 0x6F)         # Xanh Teal
GREEN = RGBColor(0x21, 0x7A, 0x4B)        # Xanh lá tích cực
RED = RGBColor(0xB8, 0x32, 0x28)          # Đỏ cảnh báo / điểm nghẽn
INK = RGBColor(0x1E, 0x24, 0x30)          # Đen chữ chính
MUTED = RGBColor(0x5A, 0x64, 0x78)        # Xám ghi chữ phụ
LIGHT = RGBColor(0xF8, 0xFA, 0xFC)        # Nền thẻ sáng
LINE = RGBColor(0xCB, 0xD5, 0xE1)         # Đường viền ngăn cách
WHITE = RGBColor(0xFF, 0xFF, 0xFF)        # Trắng tinh

FONT = "Calibri"
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def _add_run(paragraph, text=""):
    r = paragraph.add_run()
    if text:
        r.text = text
    return r


def _style_run(run, size, bold, color, italic=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color


def _header(slide, title_text, category_text):
    cat_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(12.1), Inches(0.3))
    tf_cat = cat_box.text_frame
    tf_cat.word_wrap = True
    tf_cat.margin_left = tf_cat.margin_top = tf_cat.margin_right = tf_cat.margin_bottom = 0
    p_cat = tf_cat.paragraphs[0]
    r_cat = _add_run(p_cat, category_text.upper())
    _style_run(r_cat, 11, True, AMBER)

    t_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.68), Inches(12.1), Inches(0.6))
    tf_t = t_box.text_frame
    tf_t.word_wrap = True
    tf_t.margin_left = tf_t.margin_top = tf_t.margin_right = tf_t.margin_bottom = 0
    p_t = tf_t.paragraphs[0]
    r_t = _add_run(p_t, title_text)
    _style_run(r_t, 21, True, INDIGO)

    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(1.32), Inches(12.133), Pt(1.5))
    line.fill.solid()
    line.fill.fore_color.rgb = LINE
    line.line.color.rgb = LINE


def _footer(slide, current_page, total_pages=10):
    foot_line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(7.05), Inches(12.133), Pt(1))
    foot_line.fill.solid()
    foot_line.fill.fore_color.rgb = LINE
    foot_line.line.color.rgb = LINE

    tb = slide.shapes.add_textbox(Inches(0.6), Inches(7.12), Inches(9.0), Inches(0.3))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    r = _add_run(p, "CNTT-KLCN142 — Đề xuất thực đơn dinh dưỡng cá nhân hóa bằng DBO/IDBO | Tổng hợp Tuần 9 – 10 – 11")
    _style_run(r, 9.5, False, MUTED)

    tb2 = slide.shapes.add_textbox(Inches(10.5), Inches(7.12), Inches(2.2), Inches(0.3))
    tf2 = tb2.text_frame
    p2 = tf2.paragraphs[0]
    p2.alignment = PP_ALIGN.RIGHT
    r2 = _add_run(p2, f"{current_page} / {total_pages}")
    _style_run(r2, 9.5, True, MUTED)


def _card(slide, left, top, width, height, bg_color=LIGHT, border_color=LINE):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    if border_color:
        card.line.color.rgb = border_color
        card.line.width = Pt(1)
    else:
        card.line.fill.background()
    return card


def _fit_image(path: Path, box_w: float, box_h: float) -> Tuple[int, int]:
    with Image.open(path) as im:
        iw, ih = im.size
    box_aspect = box_w / box_h
    aspect = iw / ih
    if aspect > box_aspect:
        w = int(box_w)
        h = int(box_w / aspect)
    else:
        h = int(box_h)
        w = int(box_h * aspect)
    return w, h


# =========================================================================== #
# SLIDE 1: BÌA BÁO CÁO TỔNG HỢP TUẦN 9 - 10 - 11
# =========================================================================== #
def build_slide_1(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = INDIGO
    bg.line.fill.background()

    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.2), Inches(0.15), Inches(4.8))
    bar.fill.solid()
    bar.fill.fore_color.rgb = AMBER
    bar.line.fill.background()

    tb = slide.shapes.add_textbox(Inches(1.2), Inches(1.2), Inches(11.0), Inches(5.0))
    tf = tb.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    r0 = _add_run(p0, "TRƯỜNG ĐẠI HỌC CÔNG THƯƠNG TP. HỒ CHÍ MINH — KHOA CÔNG NGHỆ THÔNG TIN\nKHÓA LUẬN TỐT NGHIỆP | MÃ ĐỀ TÀI: CNTT-KLCN142\n")
    _style_run(r0, 13, True, AMBER)

    p1 = tf.add_paragraph()
    p1.space_before = Pt(14)
    r1 = _add_run(p1, "Tối Ưu Hóa Thực Đơn Dinh Dưỡng Cá Nhân Hóa\nBằng Thuật Toán DBO và IDBO")
    _style_run(r1, 30, True, WHITE)

    p2 = tf.add_paragraph()
    p2.space_before = Pt(14)
    r2 = _add_run(p2, "BÁO CÁO TIẾN ĐỘ TỔNG HỢP GIAI ĐOẠN TUẦN 9 – TUẦN 10 – TUẦN 11\n(LẬP KẾ HOẠCH 7 NGÀY · HỆ THỐNG WEB API · CHUẨN HÓA BUDGET-MATCHED)")
    _style_run(r2, 16, True, RGBColor(0x6E, 0xC5, 0xC5))

    p3 = tf.add_paragraph()
    p3.space_before = Pt(28)
    r3 = _add_run(p3, "Sinh viên thực hiện:\n"
                      "• Lê Quang Duy (2001210356) — Nhóm trưởng (Thống kê, Báo cáo & Phân tích)\n"
                      "• Hồ Trung Cương (2001215682) — Phụ trách Budget-Matched & Dual Threshold\n"
                      "• Đặng Nguyễn Minh Đăng (2001230175) — Phụ trách Web/API & Trải nghiệm Người dùng\n\n"
                      "Thời gian: Tháng 10/2026 | Bộ môn: Kỹ thuật Dữ liệu & Trí tuệ Nhân tạo")
    _style_run(r3, 13, False, RGBColor(0xE2, 0xE8, 0xF0))


# =========================================================================== #
# SLIDE 2: [TUẦN 9] LẬP THỰC ĐƠN 7 NGÀY & DIVERSITY PENALTY (W9-F1)
# =========================================================================== #
def build_slide_2(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _header(slide, "[Tuần 9] Lập Thực Đơn Chu Kỳ 7 Ngày & Cơ Chế Diversity Penalty", "Kế Hoạch Dài Hạn (Tuần 9)")

    f1_path = W9_DIR / "fig_heatmap_nutrition_7days.png"
    if f1_path.exists():
        w, h = _fit_image(f1_path, Inches(7.0), Inches(5.4))
        slide.shapes.add_picture(str(f1_path), Inches(0.6), Inches(1.5), w, h)

    _card(slide, Inches(7.8), Inches(1.5), Inches(4.9), Inches(5.4))
    tb = slide.shapes.add_textbox(Inches(8.05), Inches(1.7), Inches(4.4), Inches(5.0))
    tf = tb.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    r0 = _add_run(p0, "Kỹ Thuật Thực Hiện (Tuần 9)")
    _style_run(r0, 16, True, INDIGO)

    items = [
        ("Mục tiêu mở rộng:", "Chuyển đổi từ đề xuất 1 ngày sang chu kỳ 7 ngày liên tục, giải quyết triệt để bài toán lặp món ăn gây nhàm chán."),
        ("Cơ chế Diversity Penalty:", "Tính toán tỷ lệ trùng lặp món chính giữa các ngày. Phạt nặng nếu tỷ lệ món chính độc nhất dưới ngưỡng 80%."),
        ("Độ bao phủ thực nghiệm (W9-F1):", "Heatmap 7 ngày x 6 dưỡng chất trên Profile P1 cho thấy Năng lượng bám sát mục tiêu hàng ngày (lệch chỉ 2.37%)."),
        ("Tối ưu hóa khẩu phần 5g:", "Tất cả các món trong 7 ngày đều tuân thủ làm tròn khẩu phần bước 5g, dinh dưỡng được thẩm định sau khi làm tròn."),
        ("Kết quả đạt được:", "Tỷ lệ món chính độc nhất đạt 100% trên các phương án tối ưu, không có hiện tượng lặp món chính trong tuần.")
    ]
    for title, desc in items:
        p = tf.add_paragraph()
        p.space_before = Pt(7)
        rt = _add_run(p, f"• {title} ")
        _style_run(rt, 11.5, True, INDIGO2)
        rd = _add_run(p, desc)
        _style_run(rd, 10.5, False, INK)

    _footer(slide, 2)


# =========================================================================== #
# SLIDE 3: [TUẦN 9] ĐỐI CHUẨN 3 HỒ SƠ & TRẢI NGHIỆM UX OPTION A/B/C (W9-F2 & W9-F3)
# =========================================================================== #
def build_slide_3(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _header(slide, "[Tuần 9] Đánh Giá Trên 3 Hồ Sơ Bệnh Lý & Trải Nghiệm Option A/B/C", "Đa Dạng Hóa Hồ Sơ & UX")

    f2_path = W9_DIR / "fig_barchart_profiles_p1_p2_p3.png"
    if f2_path.exists():
        w, h = _fit_image(f2_path, Inches(6.0), Inches(2.6))
        slide.shapes.add_picture(str(f2_path), Inches(0.6), Inches(1.5), w, h)

    f3_path = W9_DIR / "fig_radar_options_abc.png"
    if f3_path.exists():
        w, h = _fit_image(f3_path, Inches(6.0), Inches(2.6))
        slide.shapes.add_picture(str(f3_path), Inches(0.6), Inches(4.3), w, h)

    _card(slide, Inches(6.8), Inches(1.5), Inches(5.9), Inches(5.4))
    tb = slide.shapes.add_textbox(Inches(7.05), Inches(1.7), Inches(5.4), Inches(5.0))
    tf = tb.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    r0 = _add_run(p0, "Số Liệu Thực Nghiệm Tuần 9 (weekly_summary.csv)")
    _style_run(r0, 16, True, INDIGO)

    points = [
        ("Mở rộng 3 hồ sơ bệnh lý thực tế:", "P1 (Nam 22t, duy trì 2492 kcal) · P2 (Nữ 45t, thừa cân, tiểu đường, giảm cân 1838 kcal) · P3 (Nam 55t, tăng huyết áp, giảm natri 1500mg)."),
        ("Đối chuẩn hiệu năng 7 ngày (W9-F2):", "• P1: Mean Fitness 90.68, lệch calo 2.37%, 0.43 vi phạm/ngày.\n• P2: Mean Fitness 79.27, lệch calo 5.59%, 1.09 vi phạm/ngày.\n• P3: Mean Fitness 88.09, lệch calo 1.84%, 0.60 vi phạm/ngày."),
        ("Đa dạng hóa trải nghiệm người dùng (UX):", "Xây dựng 3 phương án thực đơn A/B/C độc lập cho cùng hồ sơ:\n• Option A (seed 9001): Fitness 90.67, lệch calo 42.3 kcal, 100% món độc nhất.\n• Option B (seed 9002): Fitness 86.55, lệch calo 76.8 kcal, 85.7% món độc nhất.\n• Option C (seed 9003): Fitness 91.35, lệch calo 14.7 kcal, 100% món độc nhất (tối ưu nhất)."),
        ("Trực quan hóa đa chiều (W9-F3):", "Biểu đồ Radar so sánh trực tiếp độ cân bằng dinh dưỡng giữa 3 Option giúp người dùng dễ dàng lựa chọn theo sở thích.")
    ]
    for title, desc in points:
        p = tf.add_paragraph()
        p.space_before = Pt(6)
        rt = _add_run(p, f"★ {title}\n  ")
        _style_run(rt, 11, True, TEAL)
        rd = _add_run(p, desc)
        _style_run(rd, 10.5, False, INK)

    _footer(slide, 3)


# =========================================================================== #
# SLIDE 4: [TUẦN 10] NGHIỆM THU WEB API & WEB UI (ẢNH CHỤP MÀN HÌNH THỰC TẾ)
# =========================================================================== #
def build_slide_4(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _header(slide, "[Tuần 10] Nghiệm Thu Hệ Thống Web API & Giao Diện Người Dùng", "Nền Tảng Triển Khai (Tuần 10)")

    ui_screenshot = W11_DIR / "fig_web_ui_actual_screenshot.png"
    if ui_screenshot.exists():
        w, h = _fit_image(ui_screenshot, Inches(7.2), Inches(5.4))
        slide.shapes.add_picture(str(ui_screenshot), Inches(0.6), Inches(1.5), w, h)

    _card(slide, Inches(8.0), Inches(1.5), Inches(4.7), Inches(5.4))
    tb = slide.shapes.add_textbox(Inches(8.25), Inches(1.7), Inches(4.2), Inches(5.0))
    tf = tb.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    r0 = _add_run(p0, "Nghiệm Thu Kỹ Thuật (PR #28)")
    _style_run(r0, 16, True, INDIGO)

    items = [
        ("Công nghệ sử dụng:", "FastAPI, Pydantic v2, Uvicorn server, Chart.js, Vanilla HTML5/CSS3/ES6 (Dark Theme Glassmorphism)."),
        ("Kiến trúc Single-Port:", "Phục vụ đồng thời REST API (/api/*) và static frontend trên cổng 8000, loại bỏ hoàn toàn lỗi CORS và cấu hình proxy."),
        ("6 Endpoints chuẩn hóa:", "GET /health, GET /profiles, POST /targets (tính BMR/TDEE), POST /menu/generate (chạy IDBO 2 tầng), GET /foods, POST /menu/swap (stub 501 minh bạch)."),
        ("Xử lý phi nghẽn:", "two_tier_optimize() bọc qua asyncio.to_thread, Event Loop không bị nghẽn khi tối ưu."),
        ("Quy chuẩn khẩu phần 5g:", "Bảng thực đơn hiển thị minh bạch Gram gốc và Gram làm tròn 5g (portion_g_rounded), tính lại dinh dưỡng chuẩn."),
        ("Trực quan hóa đa chiều:", "Biểu đồ Radar 6 trục dinh dưỡng và thanh tiến độ theo dải dung sai chuẩn (80%–120%).")
    ]
    for title, desc in items:
        p = tf.add_paragraph()
        p.space_before = Pt(6)
        rt = _add_run(p, f"• {title} ")
        _style_run(rt, 11.5, True, INDIGO2)
        rd = _add_run(p, desc)
        _style_run(rd, 10.5, False, INK)

    _footer(slide, 4)


# =========================================================================== #
# SLIDE 5: [TUẦN 11] ĐỐI ĐẦU DBO vs IDBO TRÊN TẬP ĐỘC LẬP (W11-F1)
# =========================================================================== #
def build_slide_5(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _header(slide, "[Tuần 11] Thực Nghiệm Đối Đầu DBO vs IDBO Trên Tập Độc Lập", "Chuẩn Hóa Dữ Liệu (Tuần 11)")

    f1_path = W11_DIR / "fig_fitness_comparison_seed8003_8010.png"
    if f1_path.exists():
        w, h = _fit_image(f1_path, Inches(6.5), Inches(5.4))
        slide.shapes.add_picture(str(f1_path), Inches(0.6), Inches(1.5), w, h)

    _card(slide, Inches(7.3), Inches(1.5), Inches(5.4), Inches(5.4))
    tb = slide.shapes.add_textbox(Inches(7.55), Inches(1.7), Inches(4.9), Inches(5.0))
    tf = tb.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    r0 = _add_run(p0, "Số Liệu Thực Nghiệm (Ngưỡng 0.28)")
    _style_run(r0, 16, True, INDIGO)

    points = [
        ("Dữ liệu thử nghiệm:", "Tập đánh giá độc lập gồm 8 hạt giống (seed 8003–8010), cấu hình chuẩn 100 vòng ngoài, 50 vòng trong, quần thể 20/15 cá thể."),
        ("Tỷ lệ đối đầu:", "IDBO Thắng 5/8 hạt giống (62.5%), DBO Thắng 3/8 hạt giống (37.5%)."),
        ("Điểm Fitness trung bình:", "DBO: 98.9931 ± 0.7918 | IDBO: 99.1590 ± 0.6851. IDBO cải thiện nhẹ trung bình +0.1659 điểm."),
        ("Kiểm định Wilcoxon Signed-Rank:", "Trị thống kê W = 14.0 > W_critical = 3 (n=8, alpha=0.05). Sự khác biệt chưa có ý nghĩa thống kê (p > 0.05)."),
        ("Phân định tập dữ liệu:", "Tập hiệu chỉnh (Calibration seed 8001–8002: IDBO thắng 1, DBO thắng 1) được tách riêng, không gộp vào tập đánh giá độc lập để tránh thiên lệch chọn mẫu.")
    ]
    for title, desc in points:
        p = tf.add_paragraph()
        p.space_before = Pt(8)
        rt = _add_run(p, f"• {title} ")
        _style_run(rt, 12, True, INDIGO2)
        rd = _add_run(p, desc)
        _style_run(rd, 11, False, INK)

    _footer(slide, 5)


# =========================================================================== #
# SLIDE 6: [TUẦN 11] ĐÁNH ĐỔI CHI PHÍ TÍNH TOÁN: EVALUATIONS & RUNTIME (W11-F2)
# =========================================================================== #
def build_slide_6(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _header(slide, "[Tuần 11] Đánh Đổi Tài Nguyên Tính Toán: Số Evaluations & Thời Gian Chạy", "Phân Tích Chi Phí")

    f2_path = W11_DIR / "fig_resource_tradeoff_evals_runtime.png"
    if f2_path.exists():
        w, h = _fit_image(f2_path, Inches(7.2), Inches(5.4))
        slide.shapes.add_picture(str(f2_path), Inches(0.6), Inches(1.5), w, h)

    _card(slide, Inches(8.0), Inches(1.5), Inches(4.7), Inches(5.4))
    tb = slide.shapes.add_textbox(Inches(8.25), Inches(1.7), Inches(4.2), Inches(5.0))
    tf = tb.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    r0 = _add_run(p0, "Phân Tích Chi Phí Tính Toán")
    _style_run(r0, 16, True, INDIGO)

    tradeoffs = [
        ("Số lượt tính hàm mục tiêu:", "DBO cố định ở mức 263.870 lượt. IDBO dao động từ 293.826 đến 337.808 lượt (trung bình 313.005 lượt, tăng +18.6%)."),
        ("Thời gian thực thi:", "DBO trung bình 49.39s. IDBO trung bình 56.61s (tăng +14.6% thời gian xử lý)."),
        ("Bản chất chi phí phụ trội:", "Phát sinh từ các đợt đánh giá ứng viên đa dạng (candidate diversity evaluation), đột biến Cauchy và tái khởi động quần thể khi đình trệ."),
        ("Hệ quả khoa học:", "IDBO có điểm số tốt hơn một phần do được cấp thêm ngân sách tính toán. Cần thiết phải triển khai so sánh Budget-Matched (cố định số lần tính hàm mục tiêu).")
    ]
    for t, d in tradeoffs:
        p = tf.add_paragraph()
        p.space_before = Pt(8)
        rt = _add_run(p, f"▲ {t} ")
        _style_run(rt, 12, True, AMBER)
        rd = _add_run(p, d)
        _style_run(rd, 11, False, INK)

    _footer(slide, 6)


# =========================================================================== #
# SLIDE 7: [TUẦN 11] GIẢI MÃ TELEMETRY TẦNG NGOÀI VS TẦNG TRONG (W11-F3)
# =========================================================================== #
def build_slide_7(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _header(slide, "[Tuần 11] Giải Mã Telemetry: Hoạt Động Tầng Ngoài vs Tầng Trong", "Cơ Chế IDBO")

    f3_path = W11_DIR / "fig_telemetry_activations.png"
    if f3_path.exists():
        w, h = _fit_image(f3_path, Inches(7.2), Inches(5.4))
        slide.shapes.add_picture(str(f3_path), Inches(0.6), Inches(1.5), w, h)

    _card(slide, Inches(8.0), Inches(1.5), Inches(4.7), Inches(5.4))
    tb = slide.shapes.add_textbox(Inches(8.25), Inches(1.7), Inches(4.2), Inches(5.0))
    tf = tb.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    r0 = _add_run(p0, "Giải Mã Tác Động Từng Tầng")
    _style_run(r0, 16, True, INDIGO)

    tele_pts = [
        ("Tầng ngoài (Chọn tổ hợp món):", "Kích hoạt trung bình 29.5 lần đột biến và 2.9 lần restart. Giúp quần thể thoát bế tắc ở các seed phức tạp như 8005 (69 lần) và 8009 (58 lần)."),
        ("Tầng trong (Khẩu phần gram):", "Kích hoạt cực kỳ tích cực với trung bình ~16.000 lần đột biến và 103 lần restart trên mỗi lần tối ưu hoàn chỉnh."),
        ("Hiện tượng Seed 8007 & 8008:", "Tầng ngoài hoàn toàn không kích hoạt (0 đột biến, 0 restart) nhưng IDBO vẫn cho kết quả khác biệt DBO nhờ các cơ chế phụ trợ hoạt động hiệu quả tại tầng trong."),
        ("Kết luận:", "Cải thiện của IDBO không chỉ đến từ việc chọn món mà còn đến từ khả năng tinh chỉnh khẩu phần gram liên tục ở tầng trong.")
    ]
    for t, d in tele_pts:
        p = tf.add_paragraph()
        p.space_before = Pt(8)
        rt = p.add_run()
        _style_run(rt, 12, True, TEAL)
        rt.text = f"★ {t} "
        rd = p.add_run()
        _style_run(rd, 11, False, INK)
        rd.text = d

    _footer(slide, 7)


# =========================================================================== #
# SLIDE 8: [TUẦN 11] KIỂM SOÁT NGÂN SÁCH CHUNG (BUDGET-MATCHED W11-F4)
# =========================================================================== #
def build_slide_8(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _header(slide, "[Tuần 11] Kiểm Soát Tính Công Bằng Theo Ngân Sách Tính Toán", "Budget-Matched Benchmark")

    f4_path = W11_DIR / "fig_budget_matched_smoke.png"
    if f4_path.exists():
        w, h = _fit_image(f4_path, Inches(6.5), Inches(5.4))
        slide.shapes.add_picture(str(f4_path), Inches(0.6), Inches(1.5), w, h)

    _card(slide, Inches(7.3), Inches(1.5), Inches(5.4), Inches(5.4))
    tb = slide.shapes.add_textbox(Inches(7.55), Inches(1.7), Inches(4.9), Inches(5.0))
    tf = tb.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    r0 = _add_run(p0, "Triển Khai & Kiểm Chứng (PR #30)")
    _style_run(r0, 16, True, INDIGO)

    budget_pts = [
        ("Mục tiêu kỹ thuật:", "Áp đặt giới hạn trên max_evaluations dùng chung cho cả DBO và IDBO. Dừng bộ giải ngay khi chạm trần ngân sách."),
        ("Cơ chế theo dõi xuyên suốt:", "Biến _budget_remaining kiểm soát chặt chẽ từ tầng ngoài, các vòng lặp tầng trong đến bước tối ưu cuối cùng (Final Inner)."),
        ("Khóa cấu hình & Truy vết Git:", "Script experiment_week11_budget.py tự động khóa mã băm config_hash và Git commit hash vào config_locked.json."),
        ("Kết quả Smoke Test (W11-F4):", "Thử nghiệm ngân sách 150 evaluations trên seed 8001 và 8002 cho thấy cả DBO và IDBO đều dừng chính xác ở 150/150 lượt, budget_exceeded = False."),
        ("Tính sẵn sàng:", "Sẵn sàng chạy thực nghiệm toàn diện trên tập độc lập với ngân sách chung ở Tuần 12.")
    ]
    for t, d in budget_pts:
        p = tf.add_paragraph()
        p.space_before = Pt(8)
        rt = _add_run(p, f"✓ {t} ")
        _style_run(rt, 12, True, GREEN)
        rd = _add_run(p, d)
        _style_run(rd, 11, False, INK)

    _footer(slide, 8)


# =========================================================================== #
# SLIDE 9: [TUẦN 11] KIẾN TRÚC NGƯỠNG KÉP & ĐẢM BẢO CHẤT LƯỢNG 253 TESTS PASS
# =========================================================================== #
def build_slide_9(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _header(slide, "[Tuần 11] Kiến Trúc Ngưỡng Kép & Nghiệm Thu Chất Lượng Phần Mềm", "Hoàn Thiện Kiến Trúc & Kiểm Thử")

    _card(slide, Inches(0.6), Inches(1.5), Inches(5.9), Inches(5.4))
    tb1 = slide.shapes.add_textbox(Inches(0.85), Inches(1.7), Inches(5.4), Inches(5.0))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    p1 = tf1.paragraphs[0]
    r1 = _add_run(p1, "Kiến Trúc Ngưỡng Kép (Dual Threshold)")
    _style_run(r1, 16, True, INDIGO)

    pts1 = [
        ("Phân tách bản chất không gian:", "Tầng ngoài hoạt động trên không gian nhị phân/rời rạc (chọn món); tầng trong hoạt động trên không gian liên tục (chọn số gram khẩu phần)."),
        ("Cấu trúc TwoTierConfig mới:", "• outer_diversity_threshold = 0.28 (cố định sau kiểm định độc lập).\n• inner_diversity_threshold = None (chưa tự ý gán khi chưa đo lường)."),
        ("Telemetry đo lường tầng trong:", "Thu thập inner_diversity_min_observed và inner_diversity_mean_observed làm cơ sở hiệu chỉnh."),
        ("Bảo đảm backward compatibility:", "Unit test xác nhận solver tương thích hoàn toàn khi chạy độc lập từng ngưỡng.")
    ]
    for t, d in pts1:
        p = tf1.add_paragraph()
        p.space_before = Pt(7)
        rt = _add_run(p, f"★ {t}\n  ")
        _style_run(rt, 11.5, True, TEAL)
        rd = _add_run(p, d)
        _style_run(rd, 10.5, False, INK)

    _card(slide, Inches(6.8), Inches(1.5), Inches(5.9), Inches(5.4))
    tb2 = slide.shapes.add_textbox(Inches(7.05), Inches(1.7), Inches(5.4), Inches(5.0))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    p2 = tf2.paragraphs[0]
    r2 = _add_run(p2, "Đảm Bảo Chất Lượng (253/253 Tests PASS)")
    _style_run(r2, 16, True, INDIGO)

    pts2 = [
        ("Kiểm thử tự động 100% PASS:", "253/253 tests PASSED trong 22.59s (pytest v9.1.1, Python 3.12). Phủ kín Core Solver (82), Constraints (45), Realism (36), Schemas (42), Budget (27), API (21)."),
        ("Tính tái lập tuyệt đối (Deterministic):", "Mọi hàm ngẫu nhiên đều nhận seed tường minh, RNG cục bộ (np.random.default_rng(seed))."),
        ("Làm tròn khẩu phần 5g minh bạch:", "Nghiệm liên tục -> Làm tròn 5g -> Giải mã -> Tính lại dinh dưỡng -> Thẩm định validate_menu()."),
        ("Tuân thủ giới hạn y tế (AGENTS.md):", "Không phát ngôn 'an toàn tuyệt đối', phân định rõ ràng số liệu đo đạc và ước tính (sugar_g_estimated).")
    ]
    for t, d in pts2:
        p = tf2.add_paragraph()
        p.space_before = Pt(7)
        rt = _add_run(p, f"🛡 {t}\n  ")
        _style_run(rt, 11.5, True, GREEN)
        rd = _add_run(p, d)
        _style_run(rd, 10.5, False, INK)

    _footer(slide, 9)


# =========================================================================== #
# SLIDE 10: TỔNG KẾT GIAI ĐOẠN TUẦN 9-10-11 & ĐỊNH HƯỚNG TUẦN 12
# =========================================================================== #
def build_slide_10(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _header(slide, "Tổng Kết Chuỗi Tiến Độ Tuần 9–10–11 & Kế Hoạch Tuần 12", "Tổng Kết & Định Hướng")

    _card(slide, Inches(0.6), Inches(1.5), Inches(5.9), Inches(5.4))
    tb1 = slide.shapes.add_textbox(Inches(0.85), Inches(1.7), Inches(5.4), Inches(5.0))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    p1 = tf1.paragraphs[0]
    r1 = _add_run(p1, "Chuỗi Thành Quả Giai Đoạn Tuần 9 – 10 – 11")
    _style_run(r1, 16, True, INDIGO)

    achievements = [
        ("Tuần 9 — Lập thực đơn 7 ngày đa hồ sơ:", "Triển khai thành công chu kỳ 7 ngày trên 3 hồ sơ bệnh lý (P1, P2, P3), tích hợp cơ chế Diversity Penalty đạt 100% món chính độc nhất, cung cấp 3 Option A/B/C."),
        ("Tuần 10 — Hệ thống Web API & UI:", "Đưa solver lên nền tảng Web với kiến trúc Single-port FastAPI, 6 REST endpoints, giao diện Dark Glassmorphism, Radar chart 6 trục và bảng thực đơn 5g."),
        ("Tuần 11 — Chuẩn hóa độc lập & Budget-Matched:", "Phân định minh bạch tập độc lập (seed 8003–8010), kiểm định Wilcoxon, cài đặt cơ chế Budget-Matched khóa ngân sách chung, 253/253 tests pass.")
    ]
    for t, d in achievements:
        p = tf1.add_paragraph()
        p.space_before = Pt(8)
        rt = _add_run(p, f"✓ {t}\n  ")
        _style_run(rt, 11.5, True, GREEN)
        rd = _add_run(p, d)
        _style_run(rd, 10.5, False, INK)

    _card(slide, Inches(6.8), Inches(1.5), Inches(5.9), Inches(5.4))
    tb2 = slide.shapes.add_textbox(Inches(7.05), Inches(1.7), Inches(5.4), Inches(5.0))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    p2 = tf2.paragraphs[0]
    r2 = _add_run(p2, "Kế Hoạch Nghiên Cứu Trọng Tâm Tuần 12")
    _style_run(r2, 16, True, INDIGO)

    plans = [
        ("1. Benchmark Budget-Matched Đầy Đủ:", "Chạy benchmark toàn diện trên tập seed độc lập với ngân sách chung (ví dụ 260.000 evaluations) để trả lời dứt điểm câu hỏi về tính công bằng tài nguyên giữa DBO và IDBO."),
        ("2. Hiệu Chỉnh Ngưỡng Đa Dạng Tầng Trong:", "Sử dụng dữ liệu quan sát inner_diversity_mean để thiết lập và thử nghiệm ngưỡng inner_diversity_threshold phù hợp."),
        ("3. Kiểm Thử Đầu Cuối (E2E Web Test):", "Thực hiện tự động hóa kiểm thử luồng Người dùng -> UI Web -> API -> Two-Tier Solver -> Hiển thị thực đơn."),
        ("4. Hoàn Thiện Bản Thảo Khóa Luận Tốt Nghiệp:", "Tổng hợp toàn bộ các kết quả thực nghiệm từ Tuần 1 đến Tuần 12 vào báo cáo Word chính thức của đề tài.")
    ]
    for t, d in plans:
        p = tf2.add_paragraph()
        p.space_before = Pt(8)
        rt = _add_run(p, f"★ {t}\n  ")
        _style_run(rt, 11.5, True, AMBER)
        rd = _add_run(p, d)
        _style_run(rd, 10.5, False, INK)

    _footer(slide, 10)


def main():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    print("Building Week 9 - 10 - 11 Comprehensive Presentation Slides...")
    build_slide_1(prs)
    build_slide_2(prs)
    build_slide_3(prs)
    build_slide_4(prs)
    build_slide_5(prs)
    build_slide_6(prs)
    build_slide_7(prs)
    build_slide_8(prs)
    build_slide_9(prs)
    build_slide_10(prs)

    prs.save(str(OUT_PPTX))
    prs.save(str(OUT_PPTX_W11))
    print(f"[SUCCESS] Saved to: {OUT_PPTX}")
    print(f"[SUCCESS] Saved to: {OUT_PPTX_W11}")
    print(f"[METRICS] Total slides: {len(prs.slides)}")


if __name__ == "__main__":
    main()
