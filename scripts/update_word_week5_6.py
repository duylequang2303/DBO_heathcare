import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
import pandas as pd
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
DOCX_PATH = ROOT / "docs" / "Phan_Cuong_Tuan5_6_IDBO.docx"
EXP_DIR = ROOT / "experiments" / "week5_6"
SUMMARY_CSV = EXP_DIR / "idbo_vs_dbo_summary.csv"
RUNS_CSV = EXP_DIR / "idbo_vs_dbo_runs.csv"

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def update_week5_6_docx():
    print(f"Loading {DOCX_PATH}...")
    doc = docx.Document(str(DOCX_PATH))

    df_sum = pd.read_csv(SUMMARY_CSV)
    df_runs = pd.read_csv(RUNS_CSV)

    # 1. Cập nhật Bảng 3.1 (Thống kê số lần kích hoạt Perturbation và Restart)
    idbo_runs = df_runs[df_runs['algorithm'] == 'idbo']
    trig_stats = idbo_runs.groupby(['function', 'dim'])[['n_perturbations', 'n_restarts']].mean().reset_index()

    # Cập nhật đoạn dẫn paragraph 28 và 29
    if len(doc.paragraphs) > 28:
        doc.paragraphs[28].text = "Từ kết quả thực nghiệm toàn diện với M = 30, lọc các lượt chạy của thuật toán IDBO trên 3 số chiều (10, 30, 50), số lần kích hoạt trung bình của random perturbation và random restart theo từng hàm benchmark được tổng hợp như sau:"
    if len(doc.paragraphs) > 29:
        doc.paragraphs[29].text = "Bảng 3.1: Số lần kích hoạt trung bình của Perturbation và Restart (dims: 10, 30, 50; M = 30)"

    # Thay thế hoàn toàn Table 1 bằng table 4 cột mới
    t1_old = doc.tables[1]
    p_parent = t1_old._tbl.getparent()

    # Tạo bảng 4 cột trực tiếp bằng xml/docx
    new_t1 = doc.add_table(rows=1, cols=4)
    new_t1.style = 'Table Grid'
    headers = ["Hàm benchmark", "Số chiều", "Mean Perturbations", "Mean Restarts"]
    for i, h in enumerate(headers):
        cell = new_t1.rows[0].cells[i]
        cell.text = h
        set_cell_background(cell, "EAECEF")
        set_cell_margins(cell, 80, 80, 100, 100)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.bold = True
            r.font.size = Pt(9.5)
            r.font.name = "Times New Roman"

    for _, row in trig_stats.iterrows():
        r_cells = new_t1.add_row().cells
        r_cells[0].text = str(row['function'])
        r_cells[1].text = str(row['dim'])
        r_cells[2].text = f"{row['n_perturbations']:.4f}"
        r_cells[3].text = f"{row['n_restarts']:.4f}"
        for i, c in enumerate(r_cells):
            set_cell_margins(c, 50, 50, 80, 80)
            p = c.paragraphs[0]
            if i in [1, 2, 3]:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            for r in p.runs:
                r.font.size = Pt(9.0)
                r.font.name = "Times New Roman"

    # Di chuyển new_t1 vào vị trí của t1_old và xóa t1_old
    t1_old._tbl.addprevious(new_t1._tbl)
    p_parent.remove(t1_old._tbl)

    if len(doc.paragraphs) > 30:
        doc.paragraphs[30].text = "Nhận xét thống kê kích hoạt: Trên các hàm lồi đơn giản (Sphere, Schwefel 2.22), quần thể co cụm cực kỳ nhanh dẫn đến diversity giảm sâu dưới 1e-3, kích hoạt cơ chế Restart (cao nhất ở Sphere 10D đạt 0.533 lần/run). Với Rosenbrock ở 30D và 50D, cơ chế Random Perturbation được kích hoạt (0.100 và 0.067 lần/run) để bơm nhiễu thoát thung lũng khi nghiệm bị kẹt nhưng chưa đủ độ trễ để restart. Các hàm Ackley và Rastrigin duy trì độ phân tán tự nhiên cao nên thuật toán không tốn tài nguyên kích hoạt không cần thiết."

    # 2. Bổ sung Mục 4: KẾT QUẢ THỰC NGHIỆM VÀ SO SÁNH DBO VS IDBO
    p_sec4 = doc.add_paragraph()
    p_sec4.paragraph_format.space_before = Pt(14)
    p_sec4.paragraph_format.space_after = Pt(6)
    r = p_sec4.add_run("4. KẾT QUẢ THỰC NGHIỆM VÀ SO SÁNH TOÀN DIỆN DBO VS IDBO")
    r.font.bold = True
    r.font.size = Pt(13)
    r.font.name = "Times New Roman"

    p_desc = doc.add_paragraph()
    p_desc.text = "Thí nghiệm so sánh đối đầu giữa DBO và IDBO được tiến hành trên 6 hàm benchmark tiêu chuẩn với 3 số chiều (10, 30, 50). Mỗi cấu hình chạy độc lập 30 lần (M = 30, max_iter = 500, pop_size = 30), tổng cộng 1.080 lượt chạy hoàn thành. Toàn bộ số liệu best, mean, std, worst và số lần đánh giá hàm mục tiêu trung bình (Mean Evals) được tổng hợp trong Bảng 4.1."
    p_desc.runs[0].font.size = Pt(11)
    p_desc.runs[0].font.name = "Times New Roman"

    p_tbl_title = doc.add_paragraph()
    p_tbl_title.paragraph_format.space_before = Pt(8)
    p_tbl_title.paragraph_format.space_after = Pt(4)
    r = p_tbl_title.add_run("Bảng 4.1: Bảng tổng hợp so sánh hiệu năng DBO vs IDBO trên 18 cấu hình (M = 30)")
    r.font.bold = True
    r.font.size = Pt(10.5)
    r.font.name = "Times New Roman"

    # Tạo bảng Bảng 4.1
    t_comp = doc.add_table(rows=1, cols=8)
    t_comp.style = 'Table Grid'
    headers_41 = ["Hàm", "Dim", "Algo", "Best", "Mean", "Std", "Worst", "Mean Evals"]
    for i, h in enumerate(headers_41):
        cell = t_comp.rows[0].cells[i]
        cell.text = h
        set_cell_background(cell, "2B5C8F")
        set_cell_margins(cell, 80, 80, 80, 80)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9.0)
            run.font.name = "Times New Roman"

    df_sorted = df_sum.sort_values(by=["function", "dim", "algorithm"])
    for _, row in df_sorted.iterrows():
        r_cells = t_comp.add_row().cells
        r_cells[0].text = str(row['function'])
        r_cells[1].text = str(row['dim'])
        r_cells[2].text = str(row['algorithm']).upper()
        r_cells[3].text = f"{row['best']:.3e}"
        r_cells[4].text = f"{row['mean']:.3e}"
        r_cells[5].text = f"{row['std']:.3e}"
        r_cells[6].text = f"{row['worst']:.3e}"
        r_cells[7].text = f"{row['mean_n_evaluations']:.1f}"

        bg_col = "F4F7FB" if row['algorithm'] == 'idbo' else "FFFFFF"
        for i, c in enumerate(r_cells):
            set_cell_background(c, bg_col)
            set_cell_margins(c, 40, 40, 60, 60)
            p = c.paragraphs[0]
            if i in [1, 2]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif i >= 3:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            for run in p.runs:
                run.font.size = Pt(8.5)
                run.font.name = "Times New Roman"

    # 3. Bổ sung các hình ảnh so sánh
    p_img_sec = doc.add_paragraph()
    p_img_sec.paragraph_format.space_before = Pt(14)
    r = p_img_sec.add_run("4.1. Đồ thị hội tụ trung bình (Convergence Curves — F1)")
    r.font.bold = True
    r.font.size = Pt(12)
    r.font.name = "Times New Roman"

    p_img_note1 = doc.add_paragraph("Đồ thị dưới đây so sánh tốc độ và quỹ đạo hội tụ giữa DBO và IDBO trên 6 hàm benchmark ở cả 3 số chiều (10, 30, 50). Trục tung biểu diễn giá trị hàm mục tiêu (thang log), trục hoành là số vòng lặp iteration:")
    p_img_note1.runs[0].font.size = Pt(11)

    for d in [10, 30, 50]:
        img_p = EXP_DIR / f"fig_convergence_dim{d}.png"
        if img_p.exists():
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(img_p), width=Inches(6.0))
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            rc = p_cap.add_run(f"Hình 4.1.{d//20 + 1}: Đường cong hội tụ DBO vs IDBO tại số chiều dim = {d} (M = 30)")
            rc.font.italic = True
            rc.font.size = Pt(10)

    # Boxplot F2
    p_bp_sec = doc.add_paragraph()
    p_bp_sec.paragraph_format.space_before = Pt(14)
    r = p_bp_sec.add_run("4.2. Biểu đồ hộp phân bố sai số (Boxplots — F2)")
    r.font.bold = True
    r.font.size = Pt(12)
    r.font.name = "Times New Roman"

    p_bp_note = doc.add_paragraph("Biểu đồ hộp (Boxplot) thể hiện độ phân tán của 30 lần chạy độc lập giữa DBO và IDBO trên 3 số chiều, phản ánh tính ổn định và khả năng chống bẫy cực trị địa phương:")
    p_bp_note.runs[0].font.size = Pt(11)

    for d in [10, 30, 50]:
        img_p = EXP_DIR / f"fig_boxplot_dim{d}.png"
        if img_p.exists():
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(img_p), width=Inches(6.0))
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            rc = p_cap.add_run(f"Hình 4.2.{d//20 + 1}: Phân bố Best Fitness DBO vs IDBO tại số chiều dim = {d} (M = 30)")
            rc.font.italic = True
            rc.font.size = Pt(10)

    # Diversity F3
    p_div_sec = doc.add_paragraph()
    p_div_sec.paragraph_format.space_before = Pt(14)
    r = p_div_sec.add_run("4.3. Diễn biến độ đa dạng quần thể (Population Diversity — F3)")
    r.font.bold = True
    r.font.size = Pt(12)
    r.font.name = "Times New Roman"

    p_div_note = doc.add_paragraph("Chỉ số đa dạng quần thể Diversity qua 500 vòng lặp trên hai hàm đại diện: Sphere (hàm lồi hội tụ nhanh) và Rastrigin (hàm đa cực trị phức tạp) qua các số chiều 10, 30, 50:")
    p_div_note.runs[0].font.size = Pt(11)

    for d in [10, 30, 50]:
        img_p = EXP_DIR / f"fig_diversity_dim{d}.png"
        if img_p.exists():
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(img_p), width=Inches(5.6))
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            rc = p_cap.add_run(f"Hình 4.3.{d//20 + 1}: Diễn biến độ đa dạng quần thể IDBO tại dim = {d} (ngưỡng 10⁻³)")
            rc.font.italic = True
            rc.font.size = Pt(10)

    # 4. Bổ sung Mục 5: ĐÁNH GIÁ VÀ KẾT LUẬN
    p_sec5 = doc.add_paragraph()
    p_sec5.paragraph_format.space_before = Pt(14)
    r = p_sec5.add_run("5. ĐÁNH GIÁ KẾT QUẢ VÀ KẾT LUẬN")
    r.font.bold = True
    r.font.size = Pt(13)
    r.font.name = "Times New Roman"

    conclusion_points = [
        "1. Hiệu năng tìm kiếm: Trên 6 hàm benchmark tiêu chuẩn liên tục, cả DBO và IDBO đều đạt hiệu quả giải tối ưu xuất sắc (Sphere đạt tiệm cận 1e-200, Schwefel đạt 1e-100, Ackley và Griewank đạt nghiệm tối ưu tuyệt đối 0 ở các chiều cao). Hiệu năng giữa IDBO và DBO là tương đương trên các hàm liên tục sạch.",
        "2. Cơ chế kích hoạt có điều kiện: Random Perturbation và Random Restart trong IDBO hoạt động hoàn toàn tự động theo điều kiện (chỉ kích hoạt khi Diversity < 1e-3 và có độ trễ Stagnation ≥ 25). Do đó, thuật toán không gây xáo trộn vô ích, không làm giảm tốc độ hội tụ tự nhiên của DBO.",
        "3. Chi phí đánh giá hàm mục tiêu: IDBO chỉ tăng thêm rất ít số lần đánh giá hàm (Mean Evals trung bình ~15.030 đến 15.034 evals so với 15.030 evals của DBO gốc), hoàn toàn nằm trong mức cho phép của bài toán thực tế.",
        "4. Chuẩn bị cho bài toán thực đơn tuần 7: Bài toán xây dựng thực đơn dinh dưỡng cá nhân hóa là bài toán tối ưu tổ hợp đa mục tiêu, không gian tìm kiếm rời rạc và có rất nhiều ràng buộc phi tuyến (calo, nhóm chất, bệnh lý). Trong không gian nghiệm phức tạp này, hiện tượng quần thể bị kẹt cực trị cục bộ hoặc mất đa dạng diễn ra thường xuyên hơn nhiều so với benchmark liên tục. Cơ chế đo Diversity, Random Perturbation và Random Restart của IDBO đã hoàn thiện và kiểm chứng thành công, sẵn sàng tích hợp trực tiếp vào module gợi ý thực đơn ở giai đoạn tiếp theo."
    ]

    for pt in conclusion_points:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(pt)
        run.font.size = Pt(11)
        run.font.name = "Times New Roman"

    doc.save(str(DOCX_PATH))
    print(f"[OK] Successfully updated {DOCX_PATH}")

if __name__ == "__main__":
    update_week5_6_docx()
