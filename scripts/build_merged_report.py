import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
import pandas as pd
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
EXP_W4 = ROOT / "experiments" / "week4"
EXP_W5 = ROOT / "experiments" / "week5_6"
OUT_MERGED = ROOT / "docs" / "Bao_Cao_Tien_Do_Tuan_4_5_6_Cuong.docx"
OUT_W4 = ROOT / "docs" / "Bao_Cao_Tuan_4_Cuong.docx"

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=70, bottom=70, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="D3D3D3"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="2B5C8F"/>'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="none"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(13.5)
    r.font.bold = True
    r.font.color.rgb = RGBColor(43, 92, 143) # Deep Blue
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(11.5)
    r.font.bold = True
    r.font.color.rgb = RGBColor(30, 30, 30)
    return p

def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(9.5)
    r.font.italic = True
    r.font.color.rgb = RGBColor(80, 80, 80)
    return p

def add_bullet(doc, title, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(3)
    r_t = p.add_run(title + ": ")
    r_t.font.name = "Times New Roman"
    r_t.font.size = Pt(10.5)
    r_t.font.bold = True
    r_b = p.add_run(text)
    r_b.font.name = "Times New Roman"
    r_b.font.size = Pt(10.5)
    return p

def generate_report():
    print("Generating streamlined, results-first merged report...")
    doc = docx.Document()

    # Căn lề chuẩn văn bản: Top 2cm, Bottom 2cm, Left 2.5cm, Right 2cm
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(0.8)

    # ──────────────────────────────────────────────────────────────────────────
    # TRANG BÌA
    # ──────────────────────────────────────────────────────────────────────────
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("BỘ CÔNG THƯƠNG\nTRƯỜNG ĐẠI HỌC CÔNG THƯƠNG TP. HỒ CHÍ MINH\nKHOA CÔNG NGHỆ THÔNG TIN\n")
    r.font.name = "Times New Roman"
    r.font.size = Pt(11)
    r.font.bold = True

    p_div = doc.add_paragraph()
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_div = p_div.add_run("———————————————\n\n\n")
    r_div.font.name = "Times New Roman"

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_title.add_run("BÁO CÁO TIẾN ĐỘ THỰC NGHIỆM (TUẦN 4 – 5 – 6)\n")
    r_sub.font.name = "Times New Roman"
    r_sub.font.size = Pt(15)
    r_sub.font.bold = True
    r_sub.font.color.rgb = RGBColor(43, 92, 143)

    r_main = p_title.add_run(
        "ĐỀ TÀI: NGHIÊN CỨU VÀ XÂY DỰNG HỆ THỐNG ĐỀ XUẤT THỰC ĐƠN DINH DƯỠNG CÁ NHÂN HÓA DỰA TRÊN THUẬT TOÁN DUNG BEETLE OPTIMIZER CẢI TIẾN VỚI CƠ CHẾ NGẪU NHIÊN\n"
    )
    r_main.font.name = "Times New Roman"
    r_main.font.size = Pt(13)
    r_main.font.bold = True

    r_code = p_title.add_run("MÃ ĐỀ TÀI: CNTT-KLCN142\n\n\n")
    r_code.font.name = "Times New Roman"
    r_code.font.size = Pt(12)
    r_code.font.bold = True

    # Khung thông tin SV & GVHD
    p_info = doc.add_paragraph()
    p_info.paragraph_format.left_indent = Inches(1.5)
    r_gv = p_info.add_run("GIẢNG VIÊN HƯỚNG DẪN:\n")
    r_gv.font.name = "Times New Roman"
    r_gv.font.size = Pt(11)
    r_gv.font.bold = True
    r_gvn = p_info.add_run("       ThS. Đinh Nguyễn Trọng Nghĩa (nghiadnt@huit.edu.vn)\n\n")
    r_gvn.font.name = "Times New Roman"
    r_gvn.font.size = Pt(11)

    r_sv = p_info.add_run("SINH VIÊN THỰC HIỆN:\n")
    r_sv.font.name = "Times New Roman"
    r_sv.font.size = Pt(11)
    r_sv.font.bold = True
    r_svn = p_info.add_run(
        "       1. Lê Quang Duy       – MSSV: 2001230123 – Lớp: 14DHTH15\n"
        "       2. Đặng Nguyễn Minh Đăng – MSSV: 2001230175 – Lớp: 14DHTH15\n"
        "       3. Hồ Trung Cương     – MSSV: 2001230070 – Lớp: 14DHTH15\n\n\n\n"
    )
    r_svn.font.name = "Times New Roman"
    r_svn.font.size = Pt(11)

    p_bot = doc.add_paragraph()
    p_bot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_bot = p_bot.add_run("TP. HỒ CHÍ MINH, THÁNG 09/2026")
    r_bot.font.name = "Times New Roman"
    r_bot.font.size = Pt(11)
    r_bot.font.bold = True

    doc.add_page_break()

    # ──────────────────────────────────────────────────────────────────────────
    # MỤC 1: CẤU HÌNH THÍ NGHIỆM
    # ──────────────────────────────────────────────────────────────────────────
    add_heading_1(doc, "1. CẤU HÌNH THÍ NGHIỆM CHUẨN (M = 30, DIMS = 10, 30, 50)")

    p_lead = doc.add_paragraph("Thực nghiệm được chuẩn hóa đồng nhất trên 6 hàm benchmark kinh điển nhằm tạo baseline vững chắc cho DBO gốc (Tuần 4) và đối đầu trực tiếp với thuật toán cải tiến IDBO (Tuần 5–6). Bỏ số chiều dim=2 do không gian quá đơn giản; tập trung khảo sát các không gian đa chiều 10, 30 và 50.")
    p_lead.runs[0].font.name = "Times New Roman"
    p_lead.runs[0].font.size = Pt(10.5)

    # Bảng 1.1: Tham số cấu hình
    add_caption(doc, "Bảng 1.1: Cấu hình tham số thực nghiệm DBO và IDBO")
    t1 = doc.add_table(rows=1, cols=2)
    t1.alignment = WD_TABLE_ALIGNMENT.CENTER
    t1.rows[0].cells[0].text = "Tham số"
    t1.rows[0].cells[1].text = "Giá trị thiết lập"
    for c in t1.rows[0].cells:
        set_cell_background(c, "2B5C8F")
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
            r.font.size = Pt(9.5)

    params_data = [
        ("Kích thước quần thể (n_agents)", "30 cá thể"),
        ("Số vòng lặp tối đa (max_iter)", "500 iterations"),
        ("Số lần chạy độc lập (M)", "30 runs (độc lập seed)"),
        ("Số chiều đánh giá (dimensions)", "10, 30, 50 (bỏ dim=2)"),
        ("Số hàm benchmark", "6 hàm (Sphere, Schwefel 2.22, Rosenbrock, Rastrigin, Ackley, Griewank)"),
        ("Tổng số lượt chạy Tuần 4 (DBO)", "6 hàm × 3 dims × 30 runs = 540 runs"),
        ("Tổng số lượt chạy Tuần 5–6 (DBO vs IDBO)", "2 thuật toán × 6 hàm × 3 dims × 30 runs = 1.080 runs"),
        ("Cơ chế IDBO", "Diversity Threshold = 10⁻³, Perturbation Rate = 0.2, Stagnation Window = 25, Restart Rate = 0.25")
    ]
    for k, v in params_data:
        rc = t1.add_row().cells
        rc[0].text = k
        rc[1].text = v
        for c in rc:
            set_cell_margins(c, 50, 50, 80, 80)
            for r in c.paragraphs[0].runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(9.5)
    set_table_borders(t1)

    # Bảng 1.2: Danh mục 6 hàm benchmark
    add_caption(doc, "\nBảng 1.2: Danh mục 6 hàm benchmark tiêu chuẩn")
    t12 = doc.add_table(rows=1, cols=4)
    t12.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(["Hàm", "Phân loại", "Miền tìm kiếm [lb, ub]", "Nghiệm tối ưu toàn cục"]):
        t12.rows[0].cells[i].text = h
        set_cell_background(t12.rows[0].cells[i], "EAECEF")
        p = t12.rows[0].cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.bold = True
            r.font.size = Pt(9.5)

    bench_data = [
        ("Sphere", "Unimodal (lồi trơn)", "[-100, 100]", "x* = 0, f* = 0"),
        ("Schwefel 2.22", "Unimodal", "[-10, 10]", "x* = 0, f* = 0"),
        ("Rosenbrock", "Unimodal (thung lũng cong hẹp)", "[-30, 30]", "x* = 1, f* = 0"),
        ("Rastrigin", "Multimodal (rất nhiều cực trị địa phương)", "[-5.12, 5.12]", "x* = 0, f* = 0"),
        ("Ackley", "Multimodal (lòng chảo dốc)", "[-32, 32]", "x* = 0, f* = 0"),
        ("Griewank", "Multimodal (nhiều cực trị phụ liên kết)", "[-600, 600]", "x* = 0, f* = 0"),
    ]
    for b in bench_data:
        rc = t12.add_row().cells
        for i, val in enumerate(b):
            rc[i].text = val
            set_cell_margins(rc[i], 40, 40, 60, 60)
            for r in rc[i].paragraphs[0].runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(9.0)
    set_table_borders(t12)

    # ──────────────────────────────────────────────────────────────────────────
    # MỤC 2: KẾT QUẢ DBO GỐC (TUẦN 4)
    # ──────────────────────────────────────────────────────────────────────────
    add_heading_1(doc, "2. KẾT QUẢ THỰC NGHIỆM DBO GỐC (TUẦN 4 — M = 30)")

    add_caption(doc, "Bảng 2.1: Kết quả định lượng DBO trên 18 cấu hình thực nghiệm (dims: 10, 30, 50)")
    t21 = doc.add_table(rows=1, cols=7)
    t21.alignment = WD_TABLE_ALIGNMENT.CENTER
    h_w4 = ["Hàm", "Dim", "Runs", "Best", "Mean", "Std", "Worst"]
    for i, h in enumerate(h_w4):
        c = t21.rows[0].cells[i]
        c.text = h
        set_cell_background(c, "2B5C8F")
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
            r.font.size = Pt(9.0)

    df_w4_sum = pd.read_csv(EXP_W4 / "dbo_benchmark_summary.csv")
    for _, row in df_w4_sum.iterrows():
        rc = t21.add_row().cells
        rc[0].text = str(row["function"])
        rc[1].text = str(row["dim"])
        rc[2].text = str(int(row["runs"]))
        rc[3].text = f"{row['best']:.3e}"
        rc[4].text = f"{row['mean']:.3e}"
        rc[5].text = f"{row['std']:.3e}"
        rc[6].text = f"{row['worst']:.3e}"
        for i, c in enumerate(rc):
            set_cell_margins(c, 40, 40, 60, 60)
            p = c.paragraphs[0]
            if i in [1, 2]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif i >= 3:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(8.5)
    set_table_borders(t21)

    # Đồ thị Tuần 4
    add_heading_2(doc, "Trực quan hóa hội tụ và phân bố sai số DBO:")
    
    img_conv_m = EXP_W4 / "fig_convergence_multidim.png"
    if img_conv_m.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(img_conv_m), width=Inches(6.0))
        add_caption(doc, "Hình 2.1: Đồ thị hội tụ trung bình của DBO qua các số chiều 10, 30, 50 (M = 30)")

    img_box_m = EXP_W4 / "fig_boxplot_multidim.png"
    if img_box_m.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(img_box_m), width=Inches(6.0))
        add_caption(doc, "Hình 2.2: Biểu đồ hộp Boxplot sai số DBO qua các số chiều 10, 30, 50 (M = 30)")

    add_heading_2(doc, "Tóm tắt kết quả DBO tuần 4:")
    add_bullet(doc, "Sphere & Schwefel 2.22", "Hội tụ siêu sâu về cực tiểu toàn cục 0 (Sphere đạt tiệm cận 10⁻¹⁵⁰ ~ 10⁻²⁰⁰; Schwefel đạt 10⁻⁸² ~ 10⁻¹⁰⁶). Khả năng tìm kiếm không bị suy giảm khi tăng lên 50 chiều.")
    add_bullet(doc, "Ackley & Griewank", "Ackley đạt chính xác ngưỡng sai số dấu phẩy động 4.44×10⁻¹⁶ trên 100% các run và mọi số chiều. Griewank đạt tối ưu tuyệt đối 0 ở cả 30D và 50D.")
    add_bullet(doc, "Rosenbrock", "Là hàm thử thách nhất do đáy thung lũng hẹp và xoắn; sai số tăng tỉ lệ thuận theo số chiều (mean: 5.61 ở 10D → 26.23 ở 30D → 46.65 ở 50D).")
    add_bullet(doc, "Rastrigin", "Ở 10D và 30D còn một số lần chạy bị vướng cực trị địa phương, nhưng ở 50D đạt độ ổn định hoàn hảo với 100% các lần chạy (30/30 runs) đạt đúng giá trị 0.")

    # ──────────────────────────────────────────────────────────────────────────
    # MỤC 3: THUẬT TOÁN CẢI TIẾN IDBO & KẾT QUẢ ĐỐI ĐẦU DBO VS IDBO (TUẦN 5–6)
    # ──────────────────────────────────────────────────────────────────────────
    doc.add_page_break()
    add_heading_1(doc, "3. THUẬT TOÁN CẢI TIẾN IDBO VÀ SO SÁNH ĐỐI ĐẦU DBO VS IDBO (TUẦN 5–6)")

    p_idbo_intro = doc.add_paragraph(
        "IDBO kế thừa toàn bộ 4 hành vi sinh tồn của DBO và bổ sung cơ chế kiểm soát thích nghi sau mỗi iteration: "
        "(1) Đo Population Diversity (chuẩn hóa độ lệch chuẩn theo biên); "
        "(2) Khi Diversity < 10⁻³: Bơm nhiễu Gaussian có kiểm soát (Random Perturbation, rate=0.2, trừ elite); "
        "(3) Khi Diversity < 10⁻³ kèm Stagnation ≥ 25 vòng: Tái khởi tạo 25% cá thể kém nhất bằng phân phối đều (Random Restart) nhằm thoát bẫy cực trị địa phương."
    )
    p_idbo_intro.runs[0].font.name = "Times New Roman"
    p_idbo_intro.runs[0].font.size = Pt(10.5)

    # Bảng 3.1: Kích hoạt Perturbation & Restart
    add_caption(doc, "Bảng 3.1: Thống kê số lần kích hoạt trung bình của Perturbation và Restart trong IDBO (M = 30)")
    t31 = doc.add_table(rows=1, cols=4)
    t31.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(["Hàm benchmark", "Số chiều (Dim)", "Mean Perturbations / run", "Mean Restarts / run"]):
        t31.rows[0].cells[i].text = h
        set_cell_background(t31.rows[0].cells[i], "EAECEF")
        p = t31.rows[0].cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.bold = True
            r.font.size = Pt(9.0)

    df_runs = pd.read_csv(EXP_W5 / "idbo_vs_dbo_runs.csv")
    idbo_only = df_runs[df_runs["algorithm"] == "idbo"]
    trig_df = idbo_only.groupby(["function", "dim"])[["n_perturbations", "n_restarts"]].mean().reset_index()
    for _, row in trig_df.iterrows():
        rc = t31.add_row().cells
        rc[0].text = str(row['function'])
        rc[1].text = str(row['dim'])
        rc[2].text = f"{row['n_perturbations']:.4f}"
        rc[3].text = f"{row['n_restarts']:.4f}"
        for i, c in enumerate(rc):
            set_cell_margins(c, 30, 30, 60, 60)
            p = c.paragraphs[0]
            if i == 1:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif i >= 2:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(8.5)
    set_table_borders(t31)

    # Bảng 3.2: So sánh DBO vs IDBO
    add_caption(doc, "\nBảng 3.2: Bảng tổng hợp so sánh hiệu năng đối đầu DBO vs IDBO trên 18 cấu hình (M = 30)")
    t32 = doc.add_table(rows=1, cols=8)
    t32.alignment = WD_TABLE_ALIGNMENT.CENTER
    h_comp = ["Hàm", "Dim", "Algo", "Best", "Mean", "Std", "Worst", "Mean Evals"]
    for i, h in enumerate(h_comp):
        c = t32.rows[0].cells[i]
        c.text = h
        set_cell_background(c, "2B5C8F")
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
            r.font.size = Pt(8.5)

    df_w5_sum = pd.read_csv(EXP_W5 / "idbo_vs_dbo_summary.csv")
    df_w5_sorted = df_w5_sum.sort_values(by=["function", "dim", "algorithm"])
    for _, row in df_w5_sorted.iterrows():
        rc = t32.add_row().cells
        rc[0].text = str(row['function'])
        rc[1].text = str(row['dim'])
        rc[2].text = str(row['algorithm']).upper()
        rc[3].text = f"{row['best']:.3e}"
        rc[4].text = f"{row['mean']:.3e}"
        rc[5].text = f"{row['std']:.3e}"
        rc[6].text = f"{row['worst']:.3e}"
        rc[7].text = f"{row['mean_n_evaluations']:.1f}"

        bg_col = "F4F7FB" if row['algorithm'] == 'idbo' else "FFFFFF"
        for i, c in enumerate(rc):
            set_cell_background(c, bg_col)
            set_cell_margins(c, 30, 30, 50, 50)
            p = c.paragraphs[0]
            if i in [1, 2]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif i >= 3:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(8.0)
    set_table_borders(t32)

    # Đồ thị so sánh Tuần 5-6
    add_heading_2(doc, "3.1. Đồ thị hội tụ trung bình DBO vs IDBO (F1):")
    for d in [10, 30, 50]:
        img_p = EXP_W5 / f"fig_convergence_dim{d}.png"
        if img_p.exists():
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(img_p), width=Inches(5.8))
            add_caption(doc, f"Hình 3.1.{d//20 + 1}: Đường cong hội tụ DBO vs IDBO ở số chiều dim = {d} (M = 30)")

    add_heading_2(doc, "3.2. Biểu đồ hộp Boxplot sai số DBO vs IDBO (F2):")
    for d in [10, 30, 50]:
        img_p = EXP_W5 / f"fig_boxplot_dim{d}.png"
        if img_p.exists():
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(img_p), width=Inches(5.8))
            add_caption(doc, f"Hình 3.2.{d//20 + 1}: Phân bố sai số Best Fitness DBO vs IDBO ở dim = {d} (M = 30)")

    add_heading_2(doc, "3.3. Diễn biến độ đa dạng quần thể Diversity IDBO (F3):")
    for d in [10, 30, 50]:
        img_p = EXP_W5 / f"fig_diversity_dim{d}.png"
        if img_p.exists():
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(img_p), width=Inches(5.4))
            add_caption(doc, f"Hình 3.3.{d//20 + 1}: Độ đa dạng quần thể Diversity IDBO ở dim = {d} (Sphere vs Rastrigin)")

    add_heading_2(doc, "Tóm tắt so sánh DBO vs IDBO tuần 5–6:")
    add_bullet(doc, "Hiệu năng tương đương trên benchmark liên tục", "Trên 6 hàm liên tục toán học tiêu chuẩn, DBO gốc đã khai thác cực kỳ mạnh nên IDBO duy trì chất lượng tương đồng, không bị suy giảm nghiệm.")
    add_bullet(doc, "Cơ chế kích hoạt thông minh có điều kiện", "Perturbation và Restart chỉ can thiệp khi quần thể co cụm sâu (< 10⁻³) và bế tắc dài ngày (≥ 25 vòng). Cơ chế này kích hoạt ở Rosenbrock (bơm nhiễu 0.10 lần/run ở 30D), Sphere (restart 0.53 lần ở 10D).")
    add_bullet(doc, "Chi phí đánh giá hàm tối thiểu", "Số lần đánh giá hàm mục tiêu chỉ tăng thêm ~0.4 đến 4 evals trên tổng số 15.030 evals (Mean Evals ~15.030 đến 15.034), bảo toàn tài nguyên tính toán.")

    # ──────────────────────────────────────────────────────────────────────────
    # MỤC 4: KẾT LUẬN & ĐỊNH HƯỚNG BÀI TOÁN THỰC ĐƠN TUẦN 7
    # ──────────────────────────────────────────────────────────────────────────
    doc.add_page_break()
    add_heading_1(doc, "4. KẾT LUẬN VÀ KẾ HOẠCH BÀI TOÁN THỰC ĐƠN (TUẦN 7)")
    add_bullet(doc, "Hoàn tất kiểm thử chuẩn", "Hoàn thành 100% cả 2 giai đoạn: baseline DBO (540 runs) và so sánh đối đầu DBO vs IDBO (1.080 runs) trên 18 cấu hình, chứng minh độ ổn định và tính đúng đắn của thuật toán.")
    add_bullet(doc, "Ý nghĩa của IDBO với bài toán thực tế", "Khác với benchmark liên tục, bài toán lập thực đơn dinh dưỡng (Tuần 7) là bài toán tổ hợp rời rạc nhiều mục tiêu, không gian tìm kiếm gồ ghề và có vô số ràng buộc khắt khe (calo, protein, carb, lipid, natri, bệnh nền). Trong không gian đó, DBO gốc rất dễ sập bẫy cực trị địa phương. Cơ chế kiểm soát Diversity và Random Restart của IDBO chính là chìa khóa để thuật toán thoát bẫy và tìm ra thực đơn tối ưu.")
    add_bullet(doc, "Nhiệm vụ tuần 7", "Tích hợp thuật toán IDBO vào mô hình hóa dinh dưỡng thực tế, ánh xạ vector cá thể bọ hung sang danh sách món ăn theo bữa (Sáng - Trưa - Tối) thỏa mãn khuyến nghị RNI của Viện Dinh Dưỡng.")

    # Lưu cả 2 file: 1 file tên gộp và 1 file thay thế tuần 4
    doc.save(str(OUT_MERGED))
    print(f"[OK] Saved merged report to: {OUT_MERGED}")
    doc.save(str(OUT_W4))
    print(f"[OK] Also updated: {OUT_W4}")

if __name__ == "__main__":
    generate_report()
